"""
ElectreeFi CMS - Root Cause Analysis Word Report Generator (.docx)
Generates a structured Word report matching the exact format:
- Executive Summary with 3-cell metric banner
- Cancelled Bookings occurrence table (Counts & % Shares) + Fault Attribution table
- Low-Consumption Bookings (<1 kWh) occurrence table (Counts & % Shares) + Fault Attribution table
- Priority Recommendations
Reads directly from the generated Excel report (ElectreeFi_Cancelled_Bookings_RCA.xlsx).
"""

import os
from collections import Counter
from datetime import datetime
import openpyxl

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

# --- Color Constants ---
COLOR_NAVY = "1B365D"       # Table header & Primary Title
COLOR_LIGHT_BG = "F8FAFC"   # Alternating row background
COLOR_BORDER = "CBD5E1"     # Table borders
COLOR_CARD_BG = "F0F4F8"    # KPI card background
COLOR_MUTED = "64748B"      # Subtitles and captions
COLOR_DARK_TEXT = "1E293B"  # Body text

RGB_NAVY = RGBColor(27, 54, 93)
RGB_DARK = RGBColor(30, 41, 59)
RGB_MUTED = RGBColor(100, 116, 139)
RGB_WHITE = RGBColor(255, 255, 255)


def set_cell_shading(cell, color_hex: str):
    """Applies background fill color to a table cell."""
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shd)


def set_cell_margins(cell, top_pt=5, bottom_pt=5, left_pt=6, right_pt=6):
    """Sets internal padding for a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top_pt * 20), ('bottom', bottom_pt * 20),
                   ('left', left_pt * 20), ('right', right_pt * 20)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(int(val)))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def set_table_borders(table, color=COLOR_BORDER, sz="4"):
    """Sets subtle horizontal borders with no vertical borders."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="none"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)


# --- Comprehensive Knowledge Base for Fault Attribution ---
FAULT_ATTRIBUTION_KNOWLEDGE = {
    "remote start not received": {
        "explanation": "The connector went into Preparing state, but the charger never received a RemoteStartTransaction command from the CMS/user app before timeout.",
        "ownership": "Charger / Platform-side",
        "action": "Investigate app-to-CMS message broker latency; ensure remote start command retry logic and session watchdog are active."
    },
    "Start transaction not recieved": {
        "explanation": "The connector was in Preparing and RemoteStartTransaction was accepted by the charger, but the charger never initiated a StartTransaction.",
        "ownership": "Charger / Platform-side",
        "action": "Check charger controller communication firmware and verify digital lock / isolation pre-check sequence before start."
    },
    "gun not connected": {
        "explanation": "Booking was initiated or scheduled, but no physical vehicle coupling or OCPP log activity occurred around the booking timestamp.",
        "ownership": "User / Rider-side",
        "action": "Add in-app prompt guiding the driver to firmly insert the charging gun; send push reminder if connector remains disconnected."
    },
    "EVDisconnected": {
        "explanation": "The charging session terminated because the vehicle uncoupled or the charging gun was unplugged from the vehicle inlet.",
        "ownership": "User / Rider-side",
        "action": "Inspect connector electronic lock and mechanical latch on the gun; guide driver in app to keep gun locked until billing completes."
    },
    "Remote": {
        "explanation": "Session was terminated remotely via user app stop button or central management server command.",
        "ownership": "User-side",
        "action": "Add a short confirmation step before remote stop in app; verify driver did not stop prematurely due to perceived slow ramp-up."
    },
    "LowInsulation": {
        "explanation": "Charger insulation monitoring device (IMD) detected low insulation resistance between DC positive/negative lines and earth ground.",
        "ownership": "Grid / Site Infra",
        "action": "Conduct immediate physical insulation test on charging cable, connector pins, and contactor; inspect for moisture ingress."
    },
    "EmergencyStop": {
        "explanation": "The physical emergency stop push button was pressed, or the safety interlock circuit was tripped.",
        "ownership": "User / Site-side",
        "action": "Inspect physical E-stop switch at station; ensure reset mechanism is functioning and safety loop wiring is intact."
    },
    "Emergency Pressed": {
        "explanation": "The hardware emergency stop button was engaged on the charger facade.",
        "ownership": "User / Site-side",
        "action": "Verify if button was pressed intentionally due to site issue or accidentally; reset button and check tamper protection."
    },
    "EmergencyFault": {
        "explanation": "Emergency trip circuit or hardware safety loop triggered an emergency fault shutdown.",
        "ownership": "Site / Hardware-side",
        "action": "Perform preventative maintenance audit on charger internal safety loop, auxiliary power supply, and contactors."
    },
    "Local": {
        "explanation": "Session was stopped locally at the charger touch screen interface or via local RFID card authorization.",
        "ownership": "User-side",
        "action": "Ensure on-screen prompt clearly identifies the active connector being stopped to prevent accidental cross-gun stops."
    },
    "Local Stop": {
        "explanation": "Driver or station attendant stopped the session using the on-charger physical/screen controls.",
        "ownership": "User-side",
        "action": "Verify driver interaction UX on charger display to avoid premature session cancellation."
    },
    "PowerLoss": {
        "explanation": "Incoming AC mains utility power was interrupted or internal AC breaker tripped during charging.",
        "ownership": "Grid / Site Infra",
        "action": "Audit upstream electrical panel and transformer; check MCB/MCCB trip ratings and verify grid power stability with site host."
    },
    "AC Mains Grid Power Loss": {
        "explanation": "Utility mains supply voltage failed or fell below operating limits, causing charger shutdown.",
        "ownership": "Grid / Site Infra",
        "action": "Install auto-restart UPS for control electronics; request utility power quality audit from DISCOM."
    },
    "Earth / Ground Fault Detected": {
        "explanation": "Ground leakage current or earthing potential exceeded statutory safety threshold.",
        "ownership": "Grid / Site Infra",
        "action": "Measure site earth pit resistance (< 5 ohms standard); check neutral-to-earth voltage and cable shielding."
    },
    "Param config failed": {
        "explanation": "Digital parameter configuration handshake failed between the charger controller and vehicle BMS (e.g. voltage/current mismatch).",
        "ownership": "Charger / Platform-side",
        "action": "Update charger firmware to latest ISO 15118 / DIN 70121 compatibility profile for multi-OEM EV support."
    },
    "BMS communication error": {
        "explanation": "CAN bus or PLC communication between vehicle BMS and charger controller timed out or was interrupted.",
        "ownership": "EV / Vehicle-side",
        "action": "Inspect CP/PE communication pins in charging gun; test with vehicle diagnostic scanner to verify BMS CAN stability."
    },
    "EV Request Stop": {
        "explanation": "The vehicle on-board battery management system (BMS) transmitted an explicit stop request packet.",
        "ownership": "EV / Vehicle-side",
        "action": "Check vehicle state-of-charge limits, cell temperature alarms, and onboard charger health."
    },
    "Server Stop": {
        "explanation": "Session was halted by central management server automated policy (e.g. credit limit, billing timeout, or station watchdog).",
        "ownership": "Charger / Platform-side",
        "action": "Review backend authorization timeout logic and automated account balance enforcement rules."
    },
    "Error": {
        "explanation": "Charger controller logged an internal operating error during session handling.",
        "ownership": "Charger / Hardware-side",
        "action": "Perform controller diagnostic self-test and review vendor controller event logs."
    },
    "E_LockFail": {
        "explanation": "Electronic locking actuator failed to securely latch onto the vehicle charging inlet.",
        "ownership": "Charger / Hardware-side",
        "action": "Clean, lubricate, or replace mechanical locking pin and micro-switch actuator on the charging connector."
    },
    "Can't be specified by the logs": {
        "explanation": "Charger reported generic 'Other' stop reason with numeric status codes (e.g. 574, 575) or modem signals without granular diagnostic telemetry.",
        "ownership": "Unknown / Vendor-side",
        "action": "Coordinate with charger OEM (e.g. Tirex, Delta, Exicom) to map proprietary vendor codes to standard OCPP error descriptions."
    }
}


def get_attribution(reason: str):
    r_clean = reason.strip()
    if r_clean in FAULT_ATTRIBUTION_KNOWLEDGE:
        return FAULT_ATTRIBUTION_KNOWLEDGE[r_clean]
    low = r_clean.lower()
    if "ev" in low or "bms" in low or "vehicle" in low:
        return {
            "explanation": f"Vehicle-side or battery management interaction triggered session termination ({r_clean}).",
            "ownership": "EV / Vehicle-side",
            "action": "Verify vehicle inlet compatibility and inspect communication line integrity."
        }
    elif "gun" in low or "lock" in low:
        return {
            "explanation": f"Charging gun or mechanical connector latch issue reported ({r_clean}).",
            "ownership": "Charger / Hardware-side",
            "action": "Inspect charging gun mechanical pins and electronic locking actuator."
        }
    elif "ground" in low or "earth" in low or "grid" in low or "power" in low or "insulation" in low:
        return {
            "explanation": f"Electrical infrastructure or grid safety parameter threshold reached ({r_clean}).",
            "ownership": "Grid / Site Infra",
            "action": "Perform site electrical safety and grounding audit."
        }
    else:
        return {
            "explanation": f"Charger status event recorded as '{r_clean}' during session monitoring.",
            "ownership": "Charger / Platform-side",
            "action": "Review detailed OEM telemetry logs and update controller firmware if needed."
        }


def load_dataset_from_excel(excel_path: str = "ElectreeFi_Cancelled_Bookings_RCA.xlsx"):
    """Reads data from the RCA Excel workbook and returns aggregated statistics."""
    target_path = excel_path
    if not os.path.exists(target_path):
        target_path = "ElectreeFi_Cancelled_Bookings_RCA_All_Records.xlsx"
    if not os.path.exists(target_path):
        target_path = os.path.join("E:\\ElectreeFi", excel_path)
    if not os.path.exists(target_path):
        target_path = os.path.join("E:\\ElectreeFi", "ElectreeFi_Cancelled_Bookings_RCA_All_Records.xlsx")

    wb = openpyxl.load_workbook(target_path, data_only=True)
    ws = wb["ElectreeFi RCA 1"] if "ElectreeFi RCA 1" in wb.sheetnames else wb.active

    headers = [cell.value for cell in ws[1]]
    reason_idx = headers.index("Reason")
    rca_idx = headers.index("RCA")
    st_idx = headers.index("Station Name")

    cancelled_counts = Counter()
    completed_counts = Counter()
    stations = set()

    total_cancelled = 0
    total_completed = 0

    for row in ws.iter_rows(min_row=2, values_only=True):
        reason_type = row[reason_idx]
        rca = str(row[rca_idx]).strip()
        station = row[st_idx]
        if station:
            stations.add(station)

        if reason_type == "Cancelled":
            cancelled_counts[rca] += 1
            total_cancelled += 1
        else:
            completed_counts[rca] += 1
            total_completed += 1

    total_flagged = total_cancelled + total_completed
    return {
        "total_flagged": total_flagged,
        "total_cancelled": total_cancelled,
        "total_completed": total_completed,
        "unique_stations": len(stations),
        "cancelled_counts": cancelled_counts,
        "completed_counts": completed_counts
    }


def add_occurrence_table(doc, counts: Counter, total_count: int):
    """Builds a 3-column occurrence table: RCA Reason, Count, % Share."""
    tbl = doc.add_table(rows=len(counts) + 2, cols=3)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    tbl.columns[0].width = Inches(3.8)
    tbl.columns[1].width = Inches(1.2)
    tbl.columns[2].width = Inches(1.5)
    set_table_borders(tbl)

    # Header Row
    headers = ["RCA Reason", "Count", "% Share"]
    for idx, text in enumerate(headers):
        cell = tbl.cell(0, idx)
        set_cell_shading(cell, COLOR_NAVY)
        set_cell_margins(cell, top_pt=6, bottom_pt=6, left_pt=8, right_pt=8)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT if idx == 0 else WD_ALIGN_PARAGRAPH.CENTER if idx == 1 else WD_ALIGN_PARAGRAPH.RIGHT
        r = p.add_run(text)
        r.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGB_WHITE

    # Data Rows (sorted descending by count)
    for row_idx, (rca, count) in enumerate(counts.most_common(), start=1):
        bg = COLOR_LIGHT_BG if row_idx % 2 == 1 else "FFFFFF"
        pct = (count / total_count * 100) if total_count > 0 else 0.0

        c0 = tbl.cell(row_idx, 0)
        c1 = tbl.cell(row_idx, 1)
        c2 = tbl.cell(row_idx, 2)

        set_cell_shading(c0, bg)
        set_cell_shading(c1, bg)
        set_cell_shading(c2, bg)

        set_cell_margins(c0, top_pt=4, bottom_pt=4, left_pt=8, right_pt=8)
        set_cell_margins(c1, top_pt=4, bottom_pt=4, left_pt=8, right_pt=8)
        set_cell_margins(c2, top_pt=4, bottom_pt=4, left_pt=8, right_pt=8)

        p0 = c0.paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r0 = p0.add_run(rca)
        r0.font.size = Pt(9)
        r0.font.color.rgb = RGB_DARK

        p1 = c1.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r1 = p1.add_run(str(count))
        r1.font.size = Pt(9)
        r1.font.color.rgb = RGB_DARK

        p2 = c2.paragraphs[0]
        p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r2 = p2.add_run(f"{pct:.1f}%")
        r2.font.size = Pt(9)
        r2.font.color.rgb = RGB_DARK

    # Total Row
    tot_row = len(counts) + 1
    c_tot0 = tbl.cell(tot_row, 0)
    c_tot1 = tbl.cell(tot_row, 1)
    c_tot2 = tbl.cell(tot_row, 2)

    set_cell_shading(c_tot0, "F1F5F9")
    set_cell_shading(c_tot1, "F1F5F9")
    set_cell_shading(c_tot2, "F1F5F9")

    set_cell_margins(c_tot0, top_pt=5, bottom_pt=5, left_pt=8, right_pt=8)
    set_cell_margins(c_tot1, top_pt=5, bottom_pt=5, left_pt=8, right_pt=8)
    set_cell_margins(c_tot2, top_pt=5, bottom_pt=5, left_pt=8, right_pt=8)

    p_t0 = c_tot0.paragraphs[0]
    r_t0 = p_t0.add_run("Total")
    r_t0.bold = True
    r_t0.font.size = Pt(9.5)
    r_t0.font.color.rgb = RGB_NAVY

    p_t1 = c_tot1.paragraphs[0]
    p_t1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_t1 = p_t1.add_run(str(total_count))
    r_t1.bold = True
    r_t1.font.size = Pt(9.5)
    r_t1.font.color.rgb = RGB_NAVY

    p_t2 = c_tot2.paragraphs[0]
    p_t2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_t2 = p_t2.add_run("100%")
    r_t2.bold = True
    r_t2.font.size = Pt(9.5)
    r_t2.font.color.rgb = RGB_NAVY

    doc.add_paragraph().paragraph_format.space_after = Pt(8)


def add_attribution_table(doc, counts: Counter):
    """Builds a 4-column fault attribution table: RCA Reason, Explanation, Fault / Ownership, Recommended Action."""
    reasons = [r for r, _ in counts.most_common()]
    tbl = doc.add_table(rows=len(reasons) + 1, cols=4)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    tbl.columns[0].width = Inches(1.5)
    tbl.columns[1].width = Inches(2.3)
    tbl.columns[2].width = Inches(1.2)
    tbl.columns[3].width = Inches(1.5)
    set_table_borders(tbl)

    headers = ["RCA Reason", "Explanation", "Fault / Ownership", "Recommended Action"]
    for idx, text in enumerate(headers):
        cell = tbl.cell(0, idx)
        set_cell_shading(cell, COLOR_NAVY)
        set_cell_margins(cell, top_pt=6, bottom_pt=6, left_pt=6, right_pt=6)
        p = cell.paragraphs[0]
        r = p.add_run(text)
        r.bold = True
        r.font.size = Pt(9)
        r.font.color.rgb = RGB_WHITE

    for row_idx, rca in enumerate(reasons, start=1):
        attr = get_attribution(rca)
        bg = COLOR_LIGHT_BG if row_idx % 2 == 1 else "FFFFFF"

        c0 = tbl.cell(row_idx, 0)
        c1 = tbl.cell(row_idx, 1)
        c2 = tbl.cell(row_idx, 2)
        c3 = tbl.cell(row_idx, 3)

        set_cell_shading(c0, bg)
        set_cell_shading(c1, bg)
        set_cell_shading(c2, bg)
        set_cell_shading(c3, bg)

        set_cell_margins(c0, top_pt=4, bottom_pt=4, left_pt=6, right_pt=6)
        set_cell_margins(c1, top_pt=4, bottom_pt=4, left_pt=6, right_pt=6)
        set_cell_margins(c2, top_pt=4, bottom_pt=4, left_pt=6, right_pt=6)
        set_cell_margins(c3, top_pt=4, bottom_pt=4, left_pt=6, right_pt=6)

        p0 = c0.paragraphs[0]
        r0 = p0.add_run(rca)
        r0.bold = True
        r0.font.size = Pt(8.5)
        r0.font.color.rgb = RGB_DARK

        p1 = c1.paragraphs[0]
        r1 = p1.add_run(attr["explanation"])
        r1.font.size = Pt(8)
        r1.font.color.rgb = RGB_DARK

        p2 = c2.paragraphs[0]
        r2 = p2.add_run(attr["ownership"])
        r2.font.size = Pt(8)
        r2.font.italic = True
        r2.font.color.rgb = RGB_DARK

        p3 = c3.paragraphs[0]
        r3 = p3.add_run(attr["action"])
        r3.font.size = Pt(8)
        r3.font.color.rgb = RGB_DARK

    doc.add_paragraph().paragraph_format.space_after = Pt(12)


def build_rca_report(output_path: str = "ElectreeFi_24Hour_RCA_Report.docx") -> str:
    """Generates the full Word report matching the exact layout of the user reference screenshot."""
    data = load_dataset_from_excel()
    total_flagged = data["total_flagged"]
    total_cancelled = data["total_cancelled"]
    total_completed = data["total_completed"]
    unique_stations = data["unique_stations"]
    cancelled_counts = data["cancelled_counts"]
    completed_counts = data["completed_counts"]

    canc_pct = (total_cancelled / total_flagged * 100) if total_flagged > 0 else 0
    comp_pct = (total_completed / total_flagged * 100) if total_flagged > 0 else 0

    doc = Document()

    # 1. Page Margins (0.75" all sides)
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # 2. Document Title (Centered)
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(4)
    p_title.paragraph_format.space_after = Pt(2)
    run_title = p_title.add_run("Root Cause Analysis Report")
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = RGB_NAVY

    p_sub1 = doc.add_paragraph()
    p_sub1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub1.paragraph_format.space_before = Pt(0)
    p_sub1.paragraph_format.space_after = Pt(2)
    run_sub1 = p_sub1.add_run("Charge_IN / ElectreeFi CMS")
    run_sub1.font.name = "Calibri"
    run_sub1.font.size = Pt(12)
    run_sub1.font.bold = True
    run_sub1.font.color.rgb = RGB_NAVY

    now_str = datetime.now().strftime("%d %B %Y")
    p_sub2 = doc.add_paragraph()
    p_sub2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub2.paragraph_format.space_before = Pt(0)
    p_sub2.paragraph_format.space_after = Pt(14)
    run_sub2 = p_sub2.add_run(f"Cancelled & Low-Consumption Charging Sessions | Data as of {now_str}")
    run_sub2.font.name = "Calibri"
    run_sub2.font.size = Pt(9.5)
    run_sub2.font.italic = True
    run_sub2.font.color.rgb = RGB_MUTED

    # 3. Section 1: Executive Summary
    h1 = doc.add_paragraph()
    h1.paragraph_format.space_before = Pt(8)
    h1.paragraph_format.space_after = Pt(4)
    r_h1 = h1.add_run("1. Executive Summary")
    r_h1.font.name = "Calibri"
    r_h1.font.size = Pt(13)
    r_h1.font.bold = True
    r_h1.font.color.rgb = RGB_NAVY

    top_canc = cancelled_counts.most_common(1)[0] if cancelled_counts else ("None", 0)
    top_comp = completed_counts.most_common(1)[0] if completed_counts else ("None", 0)

    p_exec = doc.add_paragraph()
    p_exec.paragraph_format.space_before = Pt(0)
    p_exec.paragraph_format.space_after = Pt(10)
    p_exec.paragraph_format.line_spacing = 1.15
    r_exec = p_exec.add_run(
        f"Between the dataset window analysed, the network recorded {total_flagged:,} problematic bookings across "
        f"{unique_stations} stations, split into {total_cancelled:,} Cancelled bookings ({canc_pct:.1f}%) and "
        f"{total_completed:,} Low-Consumption bookings delivering under 1 kWh ({comp_pct:.1f}%). "
        f"In cancelled bookings, the leading cause was '{top_canc[0]}' accounting for {top_canc[1]} sessions "
        f"({(top_canc[1]/total_cancelled*100):.1f}% of cancelled). In completed low-consumption bookings, "
        f"'{top_comp[0]}' accounted for {top_comp[1]} sessions ({(top_comp[1]/total_completed*100):.1f}% of low-consumption). "
        f"Hardware diagnostic error test sessions have been isolated and excluded to focus exclusively on customer-facing operational health."
    )
    r_exec.font.size = Pt(9.5)
    r_exec.font.color.rgb = RGB_DARK

    # 3-Cell Metric Banner (Table)
    banner_tbl = doc.add_table(rows=1, cols=3)
    banner_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    banner_tbl.autofit = False
    for col in banner_tbl.columns:
        col.width = Inches(2.16)
    set_table_borders(banner_tbl, color=COLOR_BORDER)

    card_data = [
        (f"{total_flagged:,}", "Total Flagged Sessions"),
        (f"{total_cancelled:,} / {total_completed:,}", "Cancelled / Low-Consumption Split"),
        (f"{unique_stations}", "Stations Involved")
    ]
    for idx, (metric, label) in enumerate(card_data):
        cell = banner_tbl.cell(0, idx)
        set_cell_shading(cell, COLOR_CARD_BG)
        set_cell_margins(cell, top_pt=10, bottom_pt=10, left_pt=8, right_pt=8)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        r_num = p.add_run(f"{metric}\n")
        r_num.font.bold = True
        r_num.font.size = Pt(18)
        r_num.font.color.rgb = RGB_NAVY

        r_lbl = p.add_run(label)
        r_lbl.font.size = Pt(8.5)
        r_lbl.font.color.rgb = RGB_MUTED

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # Most Prominent Issues
    p_iss_title = doc.add_paragraph()
    p_iss_title.paragraph_format.space_before = Pt(4)
    p_iss_title.paragraph_format.space_after = Pt(4)
    r_iss_title = p_iss_title.add_run("Most Prominent Issues:")
    r_iss_title.bold = True
    r_iss_title.font.size = Pt(10)
    r_iss_title.font.color.rgb = RGB_DARK

    # Bullet points
    b1 = doc.add_paragraph(style='List Bullet')
    b1.paragraph_format.space_after = Pt(3)
    r = b1.add_run(f"{top_canc[0]} is the largest cause of session abandonment in Cancelled Bookings ({top_canc[1]} sessions, {(top_canc[1]/total_cancelled*100):.1f}%), pointing to remote start latency between user app and charger.")
    r.font.size = Pt(9)
    r.font.color.rgb = RGB_DARK

    b2 = doc.add_paragraph(style='List Bullet')
    b2.paragraph_format.space_after = Pt(3)
    r = b2.add_run(f"Premature EV Disconnections ({completed_counts.get('EVDisconnected', 0)} sessions, {(completed_counts.get('EVDisconnected', 0)/total_completed*100):.1f}% of Low-Consumption) indicate vehicle uncoupling before significant energy delivery.")
    r.font.size = Pt(9)
    r.font.color.rgb = RGB_DARK

    b3 = doc.add_paragraph(style='List Bullet')
    b3.paragraph_format.space_after = Pt(3)
    r = b3.add_run(f"Low Insulation Alarms ({completed_counts.get('LowInsulation', 0)} sessions) represent genuine infrastructure and cable isolation warnings requiring site safety audits.")
    r.font.size = Pt(9)
    r.font.color.rgb = RGB_DARK

    b4 = doc.add_paragraph(style='List Bullet')
    b4.paragraph_format.space_after = Pt(8)
    unspec_count = completed_counts.get("Can't be specified by the logs", 0)
    r = b4.add_run(f"Unspecified Vendor Codes ({unspec_count} sessions) highlight the need for OEM firmware dictionary integration.")
    r.font.size = Pt(9)
    r.font.color.rgb = RGB_DARK

    # Fault Legend
    p_leg = doc.add_paragraph()
    p_leg.paragraph_format.space_before = Pt(4)
    p_leg.paragraph_format.space_after = Pt(14)
    r_leg = p_leg.add_run("Fault Legend: User / Rider-side | Charger / Platform-side | Grid / Site Infra | Shared | Not a Fault | Unknown")
    r_leg.font.size = Pt(8)
    r_leg.font.italic = True
    r_leg.font.color.rgb = RGB_MUTED

    # 4. Section 2: Cancelled Bookings
    h2 = doc.add_paragraph()
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(4)
    r_h2 = h2.add_run("2. Cancelled Bookings")
    r_h2.font.size = Pt(13)
    r_h2.font.bold = True
    r_h2.font.color.rgb = RGB_NAVY

    p_canc_desc = doc.add_paragraph()
    p_canc_desc.paragraph_format.space_after = Pt(6)
    r_cd = p_canc_desc.add_run(
        f"Cancelled bookings are sessions that never progressed past initiation. {total_cancelled:,} such bookings "
        f"({canc_pct:.1f}% of all flagged sessions) were recorded across the network, driven primarily by remote start latency "
        f"and physical vehicle connection timeouts."
    )
    r_cd.font.size = Pt(9.5)
    r_cd.font.color.rgb = RGB_DARK

    p_occ1 = doc.add_paragraph()
    p_occ1.paragraph_format.space_before = Pt(4)
    p_occ1.paragraph_format.space_after = Pt(4)
    r_occ1 = p_occ1.add_run("Occurrence of Reasons")
    r_occ1.font.size = Pt(10.5)
    r_occ1.font.bold = True
    r_occ1.font.color.rgb = RGB_NAVY

    add_occurrence_table(doc, cancelled_counts, total_cancelled)

    p_attr1 = doc.add_paragraph()
    p_attr1.paragraph_format.space_before = Pt(4)
    p_attr1.paragraph_format.space_after = Pt(4)
    r_attr1 = p_attr1.add_run("Explanation, Fault Attribution & Recommended Fix")
    r_attr1.font.size = Pt(10.5)
    r_attr1.font.bold = True
    r_attr1.font.color.rgb = RGB_NAVY

    add_attribution_table(doc, cancelled_counts)

    # 5. Section 3: Low-Consumption Bookings (< 1 kWh)
    h3 = doc.add_paragraph()
    h3.paragraph_format.space_before = Pt(12)
    h3.paragraph_format.space_after = Pt(4)
    r_h3 = h3.add_run("3. Low-Consumption Bookings (< 1 kWh)")
    r_h3.font.size = Pt(13)
    r_h3.font.bold = True
    r_h3.font.color.rgb = RGB_NAVY

    p_comp_desc = doc.add_paragraph()
    p_comp_desc.paragraph_format.space_after = Pt(6)
    r_cpd = p_comp_desc.add_run(
        f"Low-consumption bookings are sessions that started charging but delivered under 1 kWh before ending. {total_completed:,} "
        f"such bookings ({comp_pct:.1f}% of all flagged sessions) were recorded, characterized by premature uncoupling, "
        f"remote user stops, and pre-charge safety alarms."
    )
    r_cpd.font.size = Pt(9.5)
    r_cpd.font.color.rgb = RGB_DARK

    p_occ2 = doc.add_paragraph()
    p_occ2.paragraph_format.space_before = Pt(4)
    p_occ2.paragraph_format.space_after = Pt(4)
    r_occ2 = p_occ2.add_run("Occurrence of Reasons")
    r_occ2.font.size = Pt(10.5)
    r_occ2.font.bold = True
    r_occ2.font.color.rgb = RGB_NAVY

    add_occurrence_table(doc, completed_counts, total_completed)

    p_attr2 = doc.add_paragraph()
    p_attr2.paragraph_format.space_before = Pt(4)
    p_attr2.paragraph_format.space_after = Pt(4)
    r_attr2 = p_attr2.add_run("Explanation, Fault Attribution & Recommended Fix")
    r_attr2.font.size = Pt(10.5)
    r_attr2.font.bold = True
    r_attr2.font.color.rgb = RGB_NAVY

    add_attribution_table(doc, completed_counts)

    # 6. Section 4: Priority Recommendations
    h4 = doc.add_paragraph()
    h4.paragraph_format.space_before = Pt(12)
    h4.paragraph_format.space_after = Pt(6)
    r_h4 = h4.add_run("4. Priority Recommendations")
    r_h4.font.size = Pt(13)
    r_h4.font.bold = True
    r_h4.font.color.rgb = RGB_NAVY

    recs = [
        ("Driver App Remote Stop Safeguard", "Introduce a short confirmation step before executing remote stops within the first 2 minutes of a session to prevent accidental driver stops."),
        ("Connector Locking & EV Coupling Guidance", "Inspect connector electronic lock actuators and provide visual in-app confirmation when the charging gun is properly engaged to reduce premature EVDisconnected events."),
        ("Targeted Insulation & Earthing Audits", "Conduct physical site electrical audits at stations registering repeated LowInsulation faults, checking cable integrity, earth pit resistance (< 5Ω), and contactor health."),
        ("OEM Vendor Code Standardization", "Collaborate with charger hardware manufacturers (e.g. Tirex, Delta, Exicom) to map proprietary numeric status codes (574, 575) directly to standardized OCPP alarm text in firmware."),
        ("Remote Start Watchdog Optimization", "Optimize message queue broker latency for RemoteStartTransaction to eliminate 'remote start not received' timeouts when guns are in Preparing state.")
    ]

    for title, desc in recs:
        p_rec = doc.add_paragraph(style='List Bullet')
        p_rec.paragraph_format.space_after = Pt(4)
        r_t = p_rec.add_run(f"{title}: ")
        r_t.bold = True
        r_t.font.size = Pt(9.5)
        r_t.font.color.rgb = RGB_NAVY

        r_d = p_rec.add_run(desc)
        r_d.font.size = Pt(9.5)
        r_d.font.color.rgb = RGB_DARK

    doc.save(output_path)
    return output_path


if __name__ == "__main__":
    out_file = "ElectreeFi_24Hour_RCA_Report.docx"
    build_rca_report(out_file)
    print(f"Generated Word report: {out_file}")
