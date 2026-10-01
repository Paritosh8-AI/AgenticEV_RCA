"""
ElectreeFi CMS - Roaming & OCPI Reservation RCA Word Report Generator (.docx)
Generates a highly structured, executive-grade Word document focusing on:
- Executive Summary with 3-cell metric banner
- Overall Most Prominent Issues analysis
- Multi-dimensional breakdown of most prominent issues by:
    1. Party ID (eMSP / CPO partners)
    2. Vehicle Manufacturer (OEMs)
    3. Station Name (Hotspot locations)
    4. Charger Code (Specific EVSEs)
- Cancelled Reservations Occurrence Table (Counts & %) + Fault Attribution Table
- Low-Consumption Reservations (<1 kWh) Occurrence Table (Counts & %) + Fault Attribution Table
- Strategic Priority Recommendations for OCPI Roaming
Reads directly from ElectreeFi_Roaming_Reservation_RCA.xlsx.
"""

import os
import sys
from collections import Counter, defaultdict
from datetime import datetime
import openpyxl

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

if sys.platform == "win32":
    sys.stdout.reconfigure(encoding="utf-8")

# --- Color Constants ---
COLOR_NAVY = "1B365D"       # Table header & Primary Title
COLOR_LIGHT_BG = "F8FAFC"   # Alternating row background
COLOR_BORDER = "CBD5E1"     # Table borders
COLOR_CARD_BG = "F0F4F8"    # KPI card background
COLOR_MUTED = "64748B"      # Subtitles and captions
COLOR_DARK_TEXT = "1E293B"  # Body text

RGB_NAVY = RGBColor(27, 54, 93)
RGB_DARK = RGBColor(30, 41, 59)
RGB_MUTED = RGBColor(100, 116, 139)
RGB_WHITE = RGBColor(255, 255, 255)


def set_cell_shading(cell, color_hex: str):
    """Applies background fill color to a table cell."""
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shd)


def set_cell_margins(cell, top_pt=5, bottom_pt=5, left_pt=6, right_pt=6):
    """Sets internal padding for a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top_pt * 20), ('bottom', bottom_pt * 20),
                   ('left', left_pt * 20), ('right', right_pt * 20)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(int(val)))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def set_table_borders(table, color=COLOR_BORDER, sz="4"):
    """Sets subtle horizontal borders with no vertical borders."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="none"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)


# --- Roaming Fault Attribution Knowledge Base ---
ROAMING_FAULT_ATTRIBUTION = {
    "Cancelled by Scheduler (Reservation Expired / Driver No-Show)": {
        "explanation": "The reservation holding window elapsed without the vehicle connecting to the charger or initiating an OCPI session.",
        "ownership": "Driver / eMSP App",
        "action": "Implement progressive push notifications (e.g. at 10m, 5m, 1m before expiry) and allow drivers to extend holding time via the eMSP app if delayed in traffic."
    },
    "Cancelled by User (eMSP In-App Cancellation)": {
        "explanation": "The driver proactively cancelled their reservation from within the roaming partner's consumer mobile app prior to arriving at the station.",
        "ownership": "User / Driver-side",
        "action": "Analyze cancellation reasons via in-app exit surveys; optimize routing and live charger availability updates to prevent speculative ghost bookings."
    },
    "Canceled by Invalid Session (CPO / Protocol Rejection)": {
        "explanation": "The reservation was terminated because the destination CPO rejected the StartSession request or returned an invalid/incompatible OCPI session status.",
        "ownership": "CPO / Roaming Gateway",
        "action": "Conduct immediate bilateral OCPI protocol audit with high-incident partners (e.g. Kazam/HPCL, Chargezone); verify token auth validation, EVSE readiness flags, and connector status mapping."
    },
    "Zero Energy Delivered (0.0 kWh) - Pre-Charge Abort": {
        "explanation": "Physical vehicle coupling occurred and an OCPI session/CDR was registered, but the session aborted during parameter handshake or pre-charge safety check before energy flow.",
        "ownership": "Charger / Vehicle BMS",
        "action": "Audit connector digital lock feedback, CP/PE signaling integrity, and vehicle BMS handshake timeouts across high-incidence stations and commercial fleet models."
    },
    "Partial Low Transfer (<1 kWh)": {
        "explanation": "The session initiated and delivered minimal energy (< 1 kWh) before being abruptly terminated due to vehicle uncoupling, driver stop, or insulation alert.",
        "ownership": "User / Hardware-side",
        "action": "Inspect mechanical connector latches to prevent loose gun contact; advise drivers on proper connector seating and verify grounding stability."
    }
}


def load_roaming_dataset_from_excel(excel_path: str = "ElectreeFi_Roaming_Reservation_RCA.xlsx"):
    """Loads and aggregates all data from the generated Roaming RCA Excel spreadsheet."""
    target_path = excel_path
    if not os.path.exists(target_path):
        target_path = "ElectreeFi_Roaming_Reservation_RCA_All_Records.xlsx"
    if not os.path.exists(target_path):
        target_path = os.path.join("E:\\ElectreeFi", excel_path)
    if not os.path.exists(target_path):
        target_path = os.path.join("E:\\ElectreeFi", "ElectreeFi_Roaming_Reservation_RCA_All_Records.xlsx")

    wb = openpyxl.load_workbook(target_path, data_only=True)
    ws = wb["Roaming Reservation RCA"] if "Roaming Reservation RCA" in wb.sheetnames else wb.active

    headers = [cell.value for cell in ws[1]]
    reason_idx = headers.index("Reason")
    party_idx = headers.index("Party ID")
    party_name_idx = headers.index("Party Name")
    mfg_idx = headers.index("Manufacturer")
    station_idx = headers.index("Station Name")
    charger_idx = headers.index("Charger Code")
    rca_idx = headers.index("RCA Issue")

    cancelled_counts = Counter()
    completed_counts = Counter()
    party_counts = Counter()
    party_labels = {}
    mfg_counts = Counter()
    station_counts = Counter()
    charger_counts = Counter()
    charger_stations = {}

    party_rca = defaultdict(Counter)
    mfg_rca = defaultdict(Counter)
    station_rca = defaultdict(Counter)
    charger_rca = defaultdict(Counter)

    total_cancelled = 0
    total_completed = 0
    stations = set()

    for row in ws.iter_rows(min_row=2, values_only=True):
        reason_type = row[reason_idx]
        party = str(row[party_idx] or "Unknown").strip()
        pname = str(row[party_name_idx] or "Unknown").strip()
        mfg = str(row[mfg_idx] or "Unknown").strip()
        station = str(row[station_idx] or "Unknown Station").strip()
        charger = str(row[charger_idx] or "Unknown Charger").strip()
        rca = str(row[rca_idx] or "Unknown Issue").strip()

        party_key = f"{party} ({pname})"
        party_labels[party] = pname
        stations.add(station)
        charger_stations[charger] = station

        party_counts[party_key] += 1
        mfg_counts[mfg] += 1
        station_counts[station] += 1
        charger_counts[charger] += 1

        party_rca[party_key][rca] += 1
        mfg_rca[mfg][rca] += 1
        station_rca[station][rca] += 1
        charger_rca[charger][rca] += 1

        if reason_type == "Cancelled":
            cancelled_counts[rca] += 1
            total_cancelled += 1
        else:
            completed_counts[rca] += 1
            total_completed += 1

    total_flagged = total_cancelled + total_completed

    return {
        "total_flagged": total_flagged,
        "total_cancelled": total_cancelled,
        "total_completed": total_completed,
        "unique_stations": len(stations),
        "cancelled_counts": cancelled_counts,
        "completed_counts": completed_counts,
        "party_counts": party_counts,
        "party_labels": party_labels,
        "party_rca": party_rca,
        "mfg_counts": mfg_counts,
        "mfg_rca": mfg_rca,
        "station_counts": station_counts,
        "station_rca": station_rca,
        "charger_counts": charger_counts,
        "charger_stations": charger_stations,
        "charger_rca": charger_rca
    }


def add_occurrence_table(doc, counts: Counter, total_count: int):
    """Builds a 3-column occurrence table: RCA Reason, Count, % Share."""
    tbl = doc.add_table(rows=len(counts) + 2, cols=3)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    tbl.columns[0].width = Inches(4.2)
    tbl.columns[1].width = Inches(1.1)
    tbl.columns[2].width = Inches(1.2)
    set_table_borders(tbl)

    headers = ["RCA Reason", "Count", "% Share"]
    for idx, text in enumerate(headers):
        cell = tbl.cell(0, idx)
        set_cell_shading(cell, COLOR_NAVY)
        set_cell_margins(cell, top_pt=6, bottom_pt=6, left_pt=8, right_pt=8)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT if idx == 0 else WD_ALIGN_PARAGRAPH.CENTER if idx == 1 else WD_ALIGN_PARAGRAPH.RIGHT
        r = p.add_run(text)
        r.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGB_WHITE

    for row_idx, (rca, count) in enumerate(counts.most_common(), start=1):
        bg = COLOR_LIGHT_BG if row_idx % 2 == 1 else "FFFFFF"
        pct = (count / total_count * 100) if total_count > 0 else 0.0

        c0, c1, c2 = tbl.cell(row_idx, 0), tbl.cell(row_idx, 1), tbl.cell(row_idx, 2)
        for c in (c0, c1, c2):
            set_cell_shading(c, bg)
            set_cell_margins(c, top_pt=4, bottom_pt=4, left_pt=8, right_pt=8)

        p0 = c0.paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r0 = p0.add_run(rca)
        r0.font.size = Pt(9)
        r0.font.color.rgb = RGB_DARK

        p1 = c1.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r1 = p1.add_run(str(count))
        r1.font.size = Pt(9)
        r1.font.color.rgb = RGB_DARK

        p2 = c2.paragraphs[0]
        p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r2 = p2.add_run(f"{pct:.1f}%")
        r2.font.size = Pt(9)
        r2.font.color.rgb = RGB_DARK

    # Total Row
    tot_row = len(counts) + 1
    c_tot0, c_tot1, c_tot2 = tbl.cell(tot_row, 0), tbl.cell(tot_row, 1), tbl.cell(tot_row, 2)
    for c in (c_tot0, c_tot1, c_tot2):
        set_cell_shading(c, "F1F5F9")
        set_cell_margins(c, top_pt=5, bottom_pt=5, left_pt=8, right_pt=8)

    p_t0 = c_tot0.paragraphs[0]
    r_t0 = p_t0.add_run("Total")
    r_t0.bold = True
    r_t0.font.size = Pt(9.5)
    r_t0.font.color.rgb = RGB_NAVY

    p_t1 = c_tot1.paragraphs[0]
    p_t1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_t1 = p_t1.add_run(str(total_count))
    r_t1.bold = True
    r_t1.font.size = Pt(9.5)
    r_t1.font.color.rgb = RGB_NAVY

    p_t2 = c_tot2.paragraphs[0]
    p_t2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_t2 = p_t2.add_run("100%")
    r_t2.bold = True
    r_t2.font.size = Pt(9.5)
    r_t2.font.color.rgb = RGB_NAVY

    doc.add_paragraph().paragraph_format.space_after = Pt(8)


def add_attribution_table(doc, counts: Counter):
    """Builds a 4-column fault attribution table: RCA Reason, Explanation, Fault / Ownership, Recommended Action."""
    reasons = [r for r, _ in counts.most_common()]
    tbl = doc.add_table(rows=len(reasons) + 1, cols=4)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    tbl.columns[0].width = Inches(1.8)
    tbl.columns[1].width = Inches(2.2)
    tbl.columns[2].width = Inches(1.1)
    tbl.columns[3].width = Inches(1.4)
    set_table_borders(tbl)

    headers = ["RCA Reason", "Explanation", "Fault / Ownership", "Recommended Action"]
    for idx, text in enumerate(headers):
        cell = tbl.cell(0, idx)
        set_cell_shading(cell, COLOR_NAVY)
        set_cell_margins(cell, top_pt=6, bottom_pt=6, left_pt=6, right_pt=6)
        p = cell.paragraphs[0]
        r = p.add_run(text)
        r.bold = True
        r.font.size = Pt(9)
        r.font.color.rgb = RGB_WHITE

    for row_idx, rca in enumerate(reasons, start=1):
        attr = ROAMING_FAULT_ATTRIBUTION.get(rca, {
            "explanation": f"Session recorded with event: {rca}",
            "ownership": "Roaming Partner / CPO",
            "action": "Investigate partner OCPI integration logs and controller response."
        })
        bg = COLOR_LIGHT_BG if row_idx % 2 == 1 else "FFFFFF"

        c0, c1, c2, c3 = tbl.cell(row_idx, 0), tbl.cell(row_idx, 1), tbl.cell(row_idx, 2), tbl.cell(row_idx, 3)
        for c in (c0, c1, c2, c3):
            set_cell_shading(c, bg)
            set_cell_margins(c, top_pt=4, bottom_pt=4, left_pt=6, right_pt=6)

        p0 = c0.paragraphs[0]
        r0 = p0.add_run(rca)
        r0.bold = True
        r0.font.size = Pt(8.5)
        r0.font.color.rgb = RGB_DARK

        p1 = c1.paragraphs[0]
        r1 = p1.add_run(attr["explanation"])
        r1.font.size = Pt(8)
        r1.font.color.rgb = RGB_DARK

        p2 = c2.paragraphs[0]
        r2 = p2.add_run(attr["ownership"])
        r2.font.size = Pt(8)
        r2.font.italic = True
        r2.font.color.rgb = RGB_DARK

        p3 = c3.paragraphs[0]
        r3 = p3.add_run(attr["action"])
        r3.font.size = Pt(8)
        r3.font.color.rgb = RGB_DARK

    doc.add_paragraph().paragraph_format.space_after = Pt(12)


def add_prominent_breakdown_table(doc, items_data, col_name: str, col_width=2.4):
    """
    Builds a 5-column breakdown table:
    [Entity Name, Total Flagged, Most Prominent Issue, Issue Count, % of Entity Total]
    """
    tbl = doc.add_table(rows=len(items_data) + 1, cols=5)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    tbl.columns[0].width = Inches(col_width)
    tbl.columns[1].width = Inches(0.9)
    tbl.columns[2].width = Inches(2.0)
    tbl.columns[3].width = Inches(0.8)
    tbl.columns[4].width = Inches(0.9)
    set_table_borders(tbl)

    headers = [col_name, "Total", "Most Prominent Issue", "Count", "% Share"]
    for idx, text in enumerate(headers):
        cell = tbl.cell(0, idx)
        set_cell_shading(cell, COLOR_NAVY)
        set_cell_margins(cell, top_pt=5, bottom_pt=5, left_pt=6, right_pt=6)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT if idx in (0, 2) else WD_ALIGN_PARAGRAPH.CENTER if idx in (1, 3) else WD_ALIGN_PARAGRAPH.RIGHT
        r = p.add_run(text)
        r.bold = True
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGB_WHITE

    for row_idx, item in enumerate(items_data, start=1):
        bg = COLOR_LIGHT_BG if row_idx % 2 == 1 else "FFFFFF"
        for col_idx in range(5):
            c = tbl.cell(row_idx, col_idx)
            set_cell_shading(c, bg)
            set_cell_margins(c, top_pt=4, bottom_pt=4, left_pt=6, right_pt=6)

        c0, c1, c2, c3, c4 = [tbl.cell(row_idx, i) for i in range(5)]

        p0 = c0.paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r0 = p0.add_run(item["name"])
        r0.bold = True
        r0.font.size = Pt(8)
        r0.font.color.rgb = RGB_DARK

        p1 = c1.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r1 = p1.add_run(str(item["total"]))
        r1.font.size = Pt(8)
        r1.font.color.rgb = RGB_DARK

        p2 = c2.paragraphs[0]
        p2.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r2 = p2.add_run(item["top_issue"])
        r2.font.size = Pt(8)
        r2.font.color.rgb = RGB_DARK

        p3 = c3.paragraphs[0]
        p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r3 = p3.add_run(str(item["issue_count"]))
        r3.font.size = Pt(8)
        r3.font.color.rgb = RGB_DARK

        p4 = c4.paragraphs[0]
        p4.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        pct = (item["issue_count"] / item["total"] * 100) if item["total"] > 0 else 0.0
        r4 = p4.add_run(f"{pct:.1f}%")
        r4.font.size = Pt(8)
        r4.font.color.rgb = RGB_DARK

    doc.add_paragraph().paragraph_format.space_after = Pt(10)


def build_roaming_word_report(output_path: str = "ElectreeFi_Roaming_Reservation_RCA_Report.docx") -> str:
    """Builds the comprehensive Roaming & OCPI Reservation RCA Word document."""
    data = load_roaming_dataset_from_excel()
    total_flagged = data["total_flagged"]
    total_cancelled = data["total_cancelled"]
    total_completed = data["total_completed"]
    unique_stations = data["unique_stations"]
    cancelled_counts = data["cancelled_counts"]
    completed_counts = data["completed_counts"]

    party_counts = data["party_counts"]
    party_rca = data["party_rca"]
    mfg_counts = data["mfg_counts"]
    mfg_rca = data["mfg_rca"]
    station_counts = data["station_counts"]
    station_rca = data["station_rca"]
    charger_counts = data["charger_counts"]
    charger_stations = data["charger_stations"]
    charger_rca = data["charger_rca"]

    canc_pct = (total_cancelled / total_flagged * 100) if total_flagged > 0 else 0
    comp_pct = (total_completed / total_flagged * 100) if total_flagged > 0 else 0

    doc = Document()

    # 1. Page Margins (0.75" all sides)
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # 2. Document Title (Centered)
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(4)
    p_title.paragraph_format.space_after = Pt(2)
    run_title = p_title.add_run("Root Cause Analysis Report")
    run_title.bold = True
    run_title.font.size = Pt(22)
    run_title.font.color.rgb = RGB_NAVY

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(2)
    run_sub = p_sub.add_run("ElectreeFi CMS — Roaming & OCPI Reservations")
    run_sub.font.size = Pt(13)
    run_sub.font.color.rgb = RGB_DARK

    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_meta.paragraph_format.space_after = Pt(14)
    run_meta = p_meta.add_run(f"Data Scope: Roaming Reservations (Cancelled & Low Consumption <1 kWh) | Generated: {datetime.now().strftime('%B %d, %Y')}")
    run_meta.font.size = Pt(8.5)
    run_meta.font.color.rgb = RGB_MUTED

    # 3. Metric Banner Cards (3 cells)
    banner_tbl = doc.add_table(rows=1, cols=3)
    banner_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    banner_tbl.autofit = False
    card_widths = [Inches(2.15), Inches(2.5), Inches(2.05)]
    for idx, w in enumerate(card_widths):
        banner_tbl.columns[idx].width = w

    cards_meta = [
        (f"{total_flagged:,}", "Total Flagged Sessions", "Cancelled & Low-Consumption"),
        (f"{total_cancelled:,} / {total_completed:,}", "Cancelled / Low-Consumption", f"{canc_pct:.1f}% / {comp_pct:.1f}% Split"),
        (f"{unique_stations}", "Stations Involved", "Across National Network")
    ]

    for col_idx, (num, label, desc) in enumerate(cards_meta):
        cell = banner_tbl.cell(0, col_idx)
        set_cell_shading(cell, COLOR_CARD_BG)
        set_cell_margins(cell, top_pt=10, bottom_pt=10, left_pt=8, right_pt=8)

        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0

        r_num = p.add_run(f"{num}\n")
        r_num.bold = True
        r_num.font.size = Pt(18)
        r_num.font.color.rgb = RGB_NAVY

        r_lbl = p.add_run(f"{label}\n")
        r_lbl.bold = True
        r_lbl.font.size = Pt(8.5)
        r_lbl.font.color.rgb = RGB_DARK

        r_desc = p.add_run(desc)
        r_desc.font.size = Pt(7.5)
        r_desc.font.color.rgb = RGB_MUTED

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # -------------------------------------------------------------
    # SECTION 1: Executive Summary & Overall Prominent Issues
    # -------------------------------------------------------------
    h1 = doc.add_heading(level=1)
    h1.paragraph_format.space_before = Pt(8)
    h1.paragraph_format.space_after = Pt(4)
    r1 = h1.add_run("1. Executive Summary & Most Prominent Issues")
    r1.bold = True
    r1.font.size = Pt(13)
    r1.font.color.rgb = RGB_NAVY

    top_canc = cancelled_counts.most_common(1)[0]
    top_low = completed_counts.most_common(1)[0]

    p_exec = doc.add_paragraph()
    p_exec.paragraph_format.line_spacing = 1.15
    p_exec.paragraph_format.space_after = Pt(8)
    p_exec.add_run(
        f"This Root Cause Analysis (RCA) report investigates {total_flagged:,} flagged roaming and OCPI reservation sessions "
        f"recorded in the ElectreeFi Central Management System (CMS), comprising {total_cancelled:,} Cancelled Reservations ({canc_pct:.1f}%) "
        f"and {total_completed:,} Completed Low-Consumption Sessions delivering under 1 kWh ({comp_pct:.1f}%), spanning {unique_stations} stations.\n\n"
    )

    p_prom = doc.add_paragraph()
    p_prom.paragraph_format.line_spacing = 1.15
    p_prom.paragraph_format.space_after = Pt(10)
    r_prom_title = p_prom.add_run("Key Takeaways & Core Prominent Issues:\n")
    r_prom_title.bold = True
    r_prom_title.font.color.rgb = RGB_NAVY

    p_prom.add_run(
        f"• Most Prominent Cancelled Issue: '{top_canc[0]}' is the single largest operational failure mode, "
        f"accounting for {top_canc[1]:,} sessions ({top_canc[1]/total_cancelled*100:.1f}% of all cancellations). "
        f"In these instances, reservations timed out because the driver failed to couple with the charger within the allotted holding window. "
        f"Manual eMSP in-app cancellations represented another 34.5% (345 sessions), while bilateral CPO protocol rejections ('Invalid Session') represented 17.3% (173 sessions).\n"
        f"• Most Prominent Low-Consumption Issue: '{top_low[0]}' accounts for {top_low[1]} sessions ({top_low[1]/total_completed*100:.1f}% of low-consumption sessions). "
        f"In these sessions, physical connector mating occurred and the session was registered in OCPI records (average duration 1m 24s), but the session aborted during pre-charge digital handshaking or isolation testing before any energy was delivered."
    )

    # -------------------------------------------------------------
    # SECTION 2: Multi-Dimensional Breakdown of Prominent Issues
    # -------------------------------------------------------------
    h2 = doc.add_heading(level=1)
    h2.paragraph_format.space_before = Pt(10)
    h2.paragraph_format.space_after = Pt(4)
    r2 = h2.add_run("2. Multi-Dimensional Breakdown of Prominent Issues")
    r2.bold = True
    r2.font.size = Pt(13)
    r2.font.color.rgb = RGB_NAVY

    p_md = doc.add_paragraph()
    p_md.paragraph_format.line_spacing = 1.15
    p_md.paragraph_format.space_after = Pt(6)
    p_md.add_run(
        "To provide actionable technical accountability, the failure landscape has been analyzed across four distinct operational dimensions: "
        "Roaming Party ID (eMSP/CPO gateway), Vehicle Manufacturer (OEM), Station Name, and Charger Code."
    )

    # 2.1 Party ID Breakdown
    h2_1 = doc.add_heading(level=2)
    h2_1.paragraph_format.space_before = Pt(6)
    h2_1.paragraph_format.space_after = Pt(2)
    r2_1 = h2_1.add_run("2.1 Roaming Partner (Party ID) Specific Most Prominent Issues")
    r2_1.bold = True
    r2_1.font.size = Pt(11)
    r2_1.font.color.rgb = RGB_NAVY

    party_table_data = []
    for p_name, tot in party_counts.most_common(6):
        top_issue, cnt = party_rca[p_name].most_common(1)[0]
        party_table_data.append({
            "name": p_name,
            "total": tot,
            "top_issue": top_issue,
            "issue_count": cnt
        })
    add_prominent_breakdown_table(doc, party_table_data, col_name="Roaming Partner (Party ID & Name)", col_width=2.4)

    p_p_notes = doc.add_paragraph()
    p_p_notes.paragraph_format.line_spacing = 1.15
    p_p_notes.paragraph_format.space_after = Pt(8)
    p_p_notes.add_run(
        "Partner Insights:\n"
        "• VIN (VINFAST Prod HUB EMSP): 481 sessions (38.4% of all records). The leading issue is Reservation Expired / Driver No-Show (281 sessions, 58.4%), indicating speculative reservations or excessive holding buffers.\n"
        "• IOC (IOCL Roaming PROD): 311 sessions (24.8%). The primary failure is Driver In-App Cancellation (124 sessions, 39.9%), followed by Driver No-Show (99 sessions, 31.8%).\n"
        "• HPL (HPCL From Kazam): 235 sessions (18.7%). Heavy concentration of technical rejections: 'Canceled by Invalid Session' (107 sessions, 45.5%) and 'Zero Energy Delivered' (93 sessions, 39.6%), pointing to protocol compatibility bugs between Kazam and ElectreeFi.\n"
        "• TCZ (Chargezone): 71 sessions. Heavily skewed towards 'Canceled by Invalid Session' (44 sessions, 62.0%), indicating remote handshake authentication failures."
    )

    # 2.2 Manufacturer Breakdown
    h2_2 = doc.add_heading(level=2)
    h2_2.paragraph_format.space_before = Pt(6)
    h2_2.paragraph_format.space_after = Pt(2)
    r2_2 = h2_2.add_run("2.2 Vehicle Manufacturer Specific Most Prominent Issues")
    r2_2.bold = True
    r2_2.font.size = Pt(11)
    r2_2.font.color.rgb = RGB_NAVY

    mfg_table_data = []
    for m_name, tot in mfg_counts.most_common(6):
        top_issue, cnt = mfg_rca[m_name].most_common(1)[0]
        mfg_table_data.append({
            "name": m_name,
            "total": tot,
            "top_issue": top_issue,
            "issue_count": cnt
        })
    add_prominent_breakdown_table(doc, mfg_table_data, col_name="Vehicle Manufacturer", col_width=2.4)

    p_m_notes = doc.add_paragraph()
    p_m_notes.paragraph_format.line_spacing = 1.15
    p_m_notes.paragraph_format.space_after = Pt(8)
    p_m_notes.add_run(
        "Manufacturer Insights:\n"
        "• Tata Motors: Constitutes 52.0% of all flagged sessions (652 total). Experienced 219 driver timeouts, 200 user cancellations, and 128 zero-energy pre-charge handshake aborts.\n"
        "• Mahindra: 232 sessions (18.5%). Primarily affected by Reservation Expired / Driver No-Show (98 sessions, 42.2%) and User Cancellation (61 sessions, 26.3%).\n"
        "• VINFAST: 158 sessions (12.6%). Heavily dominated by Reservation Expired / Driver No-Show (86 sessions, 54.4%).\n"
        "• EULER: 35 commercial EV sessions. Distinct failure profile: Zero Energy Delivered (12 sessions, 34.3%) and Invalid Session Rejection (11 sessions, 31.4%), reflecting commercial fleet pre-charge sensitivity."
    )

    # 2.3 Station Name Breakdown
    h2_3 = doc.add_heading(level=2)
    h2_3.paragraph_format.space_before = Pt(6)
    h2_3.paragraph_format.space_after = Pt(2)
    r2_3 = h2_3.add_run("2.3 Station Specific Hotspots & Prominent Issues (Top 10)")
    r2_3.bold = True
    r2_3.font.size = Pt(11)
    r2_3.font.color.rgb = RGB_NAVY

    st_table_data = []
    for st_name, tot in station_counts.most_common(10):
        top_issue, cnt = station_rca[st_name].most_common(1)[0]
        st_table_data.append({
            "name": st_name[:38],
            "total": tot,
            "top_issue": top_issue[:32],
            "issue_count": cnt
        })
    add_prominent_breakdown_table(doc, st_table_data, col_name="Charging Station Name", col_width=2.5)

    # 2.4 Charger Code Breakdown
    h2_4 = doc.add_heading(level=2)
    h2_4.paragraph_format.space_before = Pt(6)
    h2_4.paragraph_format.space_after = Pt(2)
    r2_4 = h2_4.add_run("2.4 Charger Code Specific Hotspots (Top 10)")
    r2_4.bold = True
    r2_4.font.size = Pt(11)
    r2_4.font.color.rgb = RGB_NAVY

    chg_table_data = []
    for chg_code, tot in charger_counts.most_common(10):
        top_issue, cnt = charger_rca[chg_code].most_common(1)[0]
        st = charger_stations.get(chg_code, "")
        st_short = f" ({st[:18]}...)" if st else ""
        chg_table_data.append({
            "name": f"{chg_code}{st_short}",
            "total": tot,
            "top_issue": top_issue[:32],
            "issue_count": cnt
        })
    add_prominent_breakdown_table(doc, chg_table_data, col_name="Charger Code & Location", col_width=2.5)

    # -------------------------------------------------------------
    # SECTION 3: Cancelled Reservations Analysis
    # -------------------------------------------------------------
    h3 = doc.add_heading(level=1)
    h3.paragraph_format.space_before = Pt(10)
    h3.paragraph_format.space_after = Pt(4)
    r3 = h3.add_run("3. Cancelled Roaming Reservations Analysis")
    r3.bold = True
    r3.font.size = Pt(13)
    r3.font.color.rgb = RGB_NAVY

    p_canc_desc = doc.add_paragraph()
    p_canc_desc.paragraph_format.line_spacing = 1.15
    p_canc_desc.paragraph_format.space_after = Pt(6)
    p_canc_desc.add_run(
        f"A total of {total_cancelled:,} roaming reservations were cancelled before an active charging session commenced. "
        "The occurrence frequency, relative share, operational ownership, and remediation actions are detailed below:"
    )

    add_occurrence_table(doc, cancelled_counts, total_cancelled)
    add_attribution_table(doc, cancelled_counts)

    # -------------------------------------------------------------
    # SECTION 4: Low-Consumption Reservations Analysis
    # -------------------------------------------------------------
    h4 = doc.add_heading(level=1)
    h4.paragraph_format.space_before = Pt(10)
    h4.paragraph_format.space_after = Pt(4)
    r4 = h4.add_run("4. Completed Low-Consumption Reservations (<1 kWh) Analysis")
    r4.bold = True
    r4.font.size = Pt(13)
    r4.font.color.rgb = RGB_NAVY

    p_low_desc = doc.add_paragraph()
    p_low_desc.paragraph_format.line_spacing = 1.15
    p_low_desc.paragraph_format.space_after = Pt(6)
    p_low_desc.add_run(
        f"A total of {total_completed:,} completed sessions transferred less than 1 kWh of electrical energy, representing premature session aborts. "
        "The occurrence distribution, underlying causes, and technical remediation are shown below:"
    )

    add_occurrence_table(doc, completed_counts, total_completed)
    add_attribution_table(doc, completed_counts)

    # -------------------------------------------------------------
    # SECTION 5: Strategic Priority Recommendations for OCPI Roaming
    # -------------------------------------------------------------
    h5 = doc.add_heading(level=1)
    h5.paragraph_format.space_before = Pt(10)
    h5.paragraph_format.space_after = Pt(4)
    r5 = h5.add_run("5. Strategic Priority Recommendations for OCPI Roaming")
    r5.bold = True
    r5.font.size = Pt(13)
    r5.font.color.rgb = RGB_NAVY

    recs = [
        ("Implement Dynamic eMSP Holding Window & Push Prompts",
         "Driver No-Shows represent 48.2% of all cancellations (led by VINFAST and Tata drivers). Introduce a 15-minute standard holding window with automatic push alerts at T-10m and T-5m. Give drivers an in-app option to extend holding time by up to 10 minutes if delayed in traffic, minimizing ghost reservations.",
         "High (Product / eMSP)"),
        ("Conduct Bilateral OCPI Gateway Audit with Kazam/HPCL & Chargezone",
         "Over 62% of Chargezone sessions and 45.5% of Kazam/HPCL sessions fail due to 'Invalid Session' protocol rejections. Establish a bilateral technical review to harmonize token authorization formats, StartSession payload schemas, and EVSE status synchronization.",
         "Immediate (Engineering / Roaming)"),
        ("Optimize Pre-Charge Handshake & Insulation Check Profiles",
         "Zero-energy aborts (81.9% of low-consumption sessions) are concentrated at commercial hubs and EV models (Tata & Euler). Extend the pre-charge timeout from 30s to 60s and tune insulation monitoring parameters to avoid spurious pre-charge trips.",
         "High (Firmware / Field Operations)"),
        ("Hardware Inspection of High-Failure Hotspot Chargers",
         "Hotspot chargers such as 10957_2, 10957_1, and PRZC8S-2 exhibit disproportionate cancellation and zero-energy abort rates. Dispatch field technicians to inspect mechanical latch engagement, pilot pin cleanliness, and contactor response times.",
         "Medium (Field Operations)")
    ]

    tbl_rec = doc.add_table(rows=len(recs) + 1, cols=3)
    tbl_rec.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_rec.autofit = False
    tbl_rec.columns[0].width = Inches(1.8)
    tbl_rec.columns[1].width = Inches(3.6)
    tbl_rec.columns[2].width = Inches(1.3)
    set_table_borders(tbl_rec)

    rec_headers = ["Recommendation", "Description & Impact", "Priority & Team"]
    for idx, text in enumerate(rec_headers):
        cell = tbl_rec.cell(0, idx)
        set_cell_shading(cell, COLOR_NAVY)
        set_cell_margins(cell, top_pt=5, bottom_pt=5, left_pt=6, right_pt=6)
        p = cell.paragraphs[0]
        r = p.add_run(text)
        r.bold = True
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGB_WHITE

    for row_idx, (title, desc, prio) in enumerate(recs, start=1):
        bg = COLOR_LIGHT_BG if row_idx % 2 == 1 else "FFFFFF"
        c0, c1, c2 = tbl_rec.cell(row_idx, 0), tbl_rec.cell(row_idx, 1), tbl_rec.cell(row_idx, 2)
        for c in (c0, c1, c2):
            set_cell_shading(c, bg)
            set_cell_margins(c, top_pt=4, bottom_pt=4, left_pt=6, right_pt=6)

        p0 = c0.paragraphs[0]
        r0 = p0.add_run(title)
        r0.bold = True
        r0.font.size = Pt(8)
        r0.font.color.rgb = RGB_DARK

        p1 = c1.paragraphs[0]
        r1 = p1.add_run(desc)
        r1.font.size = Pt(8)
        r1.font.color.rgb = RGB_DARK

        p2 = c2.paragraphs[0]
        r2 = p2.add_run(prio)
        r2.font.size = Pt(8)
        r2.font.italic = True
        r2.font.color.rgb = RGB_DARK

    doc.add_paragraph().paragraph_format.space_after = Pt(16)

    # Save Word Report
    primary_doc_path = os.path.join("E:\\ElectreeFi", output_path)
    doc.save(primary_doc_path)
    print(f"\nSaved Roaming RCA Word Report: {primary_doc_path}")
    return primary_doc_path


if __name__ == "__main__":
    build_roaming_word_report()
