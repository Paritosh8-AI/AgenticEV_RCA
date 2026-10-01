"""
ElectreeFi CMS - Cancelled Bookings RCA Excel Generator (Last 24 Hours)
Generates an Excel spreadsheet in the exact format of 'ISSUE LIST FROM CS TEAM'.
"""

import os
import sys
import asyncio
from datetime import datetime, timedelta
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from src.auth.session_manager import is_session_valid, interactive_login
from src.scraper.browser import BrowserSession
from src.scraper.bookings import fetch_bookings_list
from src.scraper.booking_details import fetch_booking_details
from src.scraper.ocpp_logs import fetch_charger_logs_beta
from src.rca.ocpp_parser import is_remote_start_status_rejected
from src.rca.engine import analyze_booking_rca


def map_to_cs_rca_text(rca_result, logs) -> str:
    """Maps internal RCA diagnostic signals to the exact CS team phrasing."""
    cause = (rca_result.root_cause or "").lower()
    cat = (rca_result.category or "").upper()
    expl = (rca_result.explanation or "").lower()

    # User rule: Only write "remote start transaction was rejected" when you find the remote start transaction status as "rejected"
    if is_remote_start_status_rejected(logs):
        return "remote start transaction was rejected"

    # StatusNotification error codes
    for e in logs:
        if e.error_code:
            err = e.error_code.lower()
            if "ground" in err:
                return "Earth / Ground Fault Detected"
            if "power" in err or "grid" in err:
                return "AC Mains Grid Power Loss"
            if "high" in err and "temp" in err:
                return "Thermal Limit Exceeded"

    if "ground" in cause or "earth" in cause:
        return "Earth / Ground Fault Detected"
    if "power loss" in cause or "grid" in cause:
        return "AC Mains Grid Power Loss"
    if "finishing" in expl or "finishing" in cause:
        return "Connector was in finishing state"
    if "suspended" in expl or "suspended" in cause:
        return "SuspendedEVSE"
    if "boot" in expl:
        return "Charger booted after remote start was received"
    if "start transaction was not received" in expl or ("aborted" in cause and "handshake" in cause):
        return "Start transaction was not received"
    if "premature vehicle disconnection" in cause or "gun was not connected" in expl or "evdisconnected" in expl:
        return "Gun was not connected"
    if "remote stop" in cause:
        return "Remote Stop triggered from user app"
    
    return rca_result.root_cause or "Gun was not connected"


async def main():
    print("=" * 70)
    print("ELECTREEFI CMS - 24-HOUR CANCELLED BOOKINGS RCA EXCEL GENERATOR")
    print("=" * 70)

    # 1. Verify Authentication Session
    if not await is_session_valid():
        print("\nSession is expired or not found.")
        print("Launching visible browser for login (CAPTCHA + OTP)...")
        res = await interactive_login(force=True)
        print(res)
        if not await is_session_valid():
            print("ERROR: Login was not completed. Please run login again.")
            sys.exit(1)

    now = datetime.now()
    start_24h = now - timedelta(hours=24)
    start_date_str = start_24h.strftime("%Y-%m-%d")
    end_date_str = now.strftime("%Y-%m-%d")

    print(f"\n[EXCEL RCA] Fetching all Cancelled bookings from {start_date_str} to {end_date_str}...")

    rows_data = []

    async with BrowserSession(headless=True) as page:
        cancelled = await fetch_bookings_list(
            page=page,
            start_date=start_date_str,
            end_date=end_date_str,
            tab="Cancelled",
            max_rows=200
        )
        print(f"Retrieved {len(cancelled)} cancelled booking candidate(s).")

        for idx, rec in enumerate(cancelled, 1):
            print(f"\n[{idx}/{len(cancelled)}] Diagnosing Booking #{rec.booking_id}...")
            try:
                details = await fetch_booking_details(
                    page=page,
                    booking_id=rec.booking_id,
                    row_selector=rec.details_link_selector
                )

                # Format Date like 9/13/2026
                raw_time = details.booking_in_time or details.scheduled_in_time or details.booking_date or rec.booking_start_time or ""
                b_dt = None
                if raw_time:
                    for fmt in [
                        "%Y-%m-%d %I:%M:%S %p", "%Y-%m-%d %I:%M %p",
                        "%d-%m-%Y %I:%M:%S %p", "%d-%m-%Y %I:%M %p",
                        "%d/%m/%Y %I:%M:%S %p", "%d/%m/%Y %I:%M %p",
                        "%d-%m-%Y %H:%M:%S", "%d/%m/%Y %H:%M:%S", "%Y-%m-%d %H:%M:%S",
                        "%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y"
                    ]:
                        try:
                            b_dt = datetime.strptime(raw_time.strip(), fmt)
                            break
                        except Exception:
                            continue

                if not b_dt:
                    b_dt = now
                b_date_str = f"{b_dt.month}/{b_dt.day}/{b_dt.year}"

                station_name = details.station_name or rec.station_organization or "Unknown Station"
                charger_code = details.charger_code or ""
                connector = details.connector_sequence_id or rec.connector_id or "1"

                # Pull logs in +/- 15 min window
                logs = []
                if charger_code:
                    w_start = (b_dt - timedelta(minutes=15)).strftime("%Y-%m-%d %H:%M")
                    w_end = (b_dt + timedelta(minutes=15)).strftime("%Y-%m-%d %H:%M")
                    logs = await fetch_charger_logs_beta(page, charger_code=charger_code, start_time=w_start, end_time=w_end)

                rca = analyze_booking_rca(booking=details, logs=logs)
                cs_rca_text = map_to_cs_rca_text(rca, logs)

                rows_data.append({
                    "date": b_date_str,
                    "reason": "Cancelled",
                    "booking_id": rec.booking_id,
                    "station_name": station_name,
                    "charger_code": charger_code,
                    "connector": str(connector),
                    "rca": cs_rca_text
                })
                print(f"  Station: {station_name}")
                print(f"  Charger: {charger_code} (Gun {connector})")
                print(f"  RCA: {cs_rca_text}")

            except Exception as e:
                print(f"  Error diagnosing {rec.booking_id}: {e}")

    # Build Excel Workbook
    print("\nGenerating Excel Workbook matching CS Team Format...")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ElectreeFi RCA 1"

    headers = ["Date", "Reason", "Booking ID", "Station Name", "Charger Code", "Connector", "RCA"]
    ws.append(headers)

    # Styling definitions matching user's reference screenshot
    yellow_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="000000")
    data_font = Font(name="Calibri", size=11, bold=False, color="000000")
    thin_border = Border(
        left=Side(style='thin', color='A0A0A0'),
        right=Side(style='thin', color='A0A0A0'),
        top=Side(style='thin', color='A0A0A0'),
        bottom=Side(style='thin', color='A0A0A0')
    )
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = yellow_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = thin_border
    ws.row_dimensions[1].height = 24

    for row_idx, r in enumerate(rows_data, start=2):
        ws.append([
            r["date"],
            r["reason"],
            int(r["booking_id"]) if r["booking_id"].isdigit() else r["booking_id"],
            r["station_name"],
            r["charger_code"],
            int(r["connector"]) if r["connector"].isdigit() else r["connector"],
            r["rca"]
        ])
        ws.row_dimensions[row_idx].height = 20
        for col_idx in range(1, len(headers) + 1):
            c = ws.cell(row=row_idx, column=col_idx)
            c.font = data_font
            c.border = thin_border
            c.alignment = center_align if col_idx in (1, 2, 3, 6) else left_align

    # Set appropriate column widths matching image
    col_widths = {
        "A": 12,  # Date
        "B": 12,  # Reason
        "C": 14,  # Booking ID
        "D": 48,  # Station Name
        "E": 22,  # Charger Code
        "F": 12,  # Connector
        "G": 58   # RCA
    }
    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width

    out_path = os.path.join("E:\\ElectreeFi", "ElectreeFi_Cancelled_Bookings_RCA.xlsx")
    wb.save(out_path)

    print("\n" + "=" * 70)
    print("SUCCESS: RCA EXCEL REPORT GENERATED!")
    print("=" * 70)
    print(f"Total Cancelled Bookings Analyzed: {len(rows_data)}")
    print(f"File Path: {out_path}")
    return out_path


if __name__ == "__main__":
    asyncio.run(main())
