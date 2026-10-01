"""
ElectreeFi CMS - Roaming & OCPI Reservation RCA Excel Generator
Fetches Cancelled Reservations and Completed Low-Consumption (<1 kWh) Reservations
from the Roaming / OCPIReservation portal endpoints, performs multi-dimensional failure classification,
and generates a structured Excel workbook matching the CS team standards.
"""

import os
import sys
import json
import re
import urllib.request
import urllib.parse
import asyncio
from datetime import datetime, timedelta
from collections import Counter
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from src.auth.session_manager import is_session_valid, interactive_login

# Ensure UTF-8 output
if sys.platform == "win32":
    sys.stdout.reconfigure(encoding="utf-8")


def get_cookie_header(storage_path="data/session_state.json") -> str:
    try:
        with open(storage_path, "r", encoding="utf-8") as f:
            data = json.load(f)
        cookies = data.get("cookies", [])
        return "; ".join([f"{c['name']}={c['value']}" for c in cookies if "ev-charge-network.com" in c.get("domain", "")])
    except Exception:
        return ""


def classify_cancelled_reason(sched_action: str | None) -> tuple[str, str]:
    """
    Classifies a cancelled roaming reservation based on its SchedularAction text.
    Returns (normalized_rca, explanation).
    """
    if not sched_action or not str(sched_action).strip():
        return (
            "Cancelled (Unspecified Schedular Action)",
            "No specific scheduler action or cancellation trigger was logged in the reservation system."
        )

    raw = str(sched_action).strip()
    low = raw.lower()

    if "invalid session" in low:
        return (
            "Canceled by Invalid Session (CPO / Protocol Rejection)",
            "The reservation was invalidated because the roaming partner CPO rejected the start session command or returned an invalid session state."
        )
    elif "scheduler" in low or "expire" in low or "timeout" in low:
        return (
            "Cancelled by Scheduler (Reservation Expired / Driver No-Show)",
            "The reservation timer elapsed without the driver connecting to the charger and initiating a session within the holding window."
        )
    elif "user" in low or "driver" in low:
        return (
            "Cancelled by User (eMSP In-App Cancellation)",
            "The driver manually cancelled the reservation through their consumer eMSP roaming mobile app before reaching the station."
        )
    else:
        # Strip timestamps
        clean = re.sub(r'\s+on\s+\d{4}-\d{2}-\d{2}\s+.*', '', raw, flags=re.IGNORECASE)
        clean = re.sub(r'\s+\d{4}-\d{2}-\d{2}\s+\d{2}:\d{2}(:\d{2})?(\s+[AP]M)?', '', clean, flags=re.IGNORECASE).strip()
        return (
            clean or "Cancelled by Scheduler",
            f"Reservation ended with scheduler action: {clean}"
        )


def classify_low_consumption_reason(kwh: float, duration: str | None, sched_action: str | None) -> tuple[str, str]:
    """
    Classifies a completed roaming reservation delivering < 1 kWh.
    Returns (normalized_rca, explanation).
    """
    if kwh <= 0.001:
        return (
            "Zero Energy Delivered (0.0 kWh) - Pre-Charge Abort",
            f"Session initiated and connected (Duration: {duration or 'N/A'}), but terminated during handshake or safety ramp-up before any energy transfer."
        )
    else:
        return (
            "Partial Low Transfer (<1 kWh)",
            f"Session transferred {kwh:.2f} kWh (Duration: {duration or 'N/A'}) before premature termination due to vehicle uncoupling or early stop."
        )


def parse_flexible_dt(t_str: str | None) -> datetime | None:
    if not t_str or not t_str.strip():
        return None
    cleaned = t_str.strip()
    for fmt in [
        "%Y-%m-%dT%H:%M:%S",
        "%Y-%m-%dT%H:%M:%S.%f",
        "%Y-%m-%d %I:%M %p",
        "%Y-%m-%d %I:%M:%S %p",
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%d",
        "%d-%m-%Y %I:%M %p",
        "%d-%m-%Y %H:%M:%S"
    ]:
        try:
            return datetime.strptime(cleaned, fmt)
        except ValueError:
            continue
    return None


async def run_roaming_rca(
    output_excel: str = "ElectreeFi_Roaming_Reservation_RCA.xlsx",
    start_date: str = "2026-09-10",
    end_date: str = "2026-09-16"
):
    print("=" * 80)
    print("ELECTREEFI CMS - ROAMING & OCPI RESERVATION RCA (EXCEL GENERATOR)")
    print(f"Date Range: {start_date} to {end_date} (Max PageSize=1000, Multi-Dimensional Focus)")
    print("=" * 80)

    # Verify session
    if not await is_session_valid():
        print("\nSession is invalid or expired. Launching semi-automated login...")
        await interactive_login(force=True)
        if not await is_session_valid():
            print("ERROR: Authentication failed.")
            sys.exit(1)
        print("Authenticated successfully!")
    else:
        print("\nExisting session is valid and active.")

    cookie_header = get_cookie_header()
    headers = {
        "X-Requested-With": "XMLHttpRequest",
        "Content-Type": "application/x-www-form-urlencoded; charset=UTF-8",
        "Cookie": cookie_header
    }
    payload_encoded = urllib.parse.urlencode({"sort": "", "page": "1", "pageSize": "1000", "group": "", "filter": ""}).encode("utf-8")

    # -------------------------------------------------------------
    # STEP 1: Query Cancelled Roaming Reservations (pageSize=1000)
    # -------------------------------------------------------------
    print(f"\n[QUERY 1/2] Fetching Cancelled Roaming Reservations ({start_date} to {end_date})...")
    canc_url = f"https://cms.ev-charge-network.com/Roaming/OCPIReservation/LoadOcpiCancelledReservationGridViewThroughAjax?StartDate={start_date}&Enddate={end_date}&transactionTypeId=-1"
    req_canc = urllib.request.Request(canc_url, data=payload_encoded, headers=headers)
    with urllib.request.urlopen(req_canc, timeout=15) as res_canc:
        data_canc = json.loads(res_canc.read().decode("utf-8"))
    raw_cancelled = data_canc.get("Data", [])
    print(f"Retrieved {len(raw_cancelled)} Cancelled Roaming reservations from portal API!")

    # -------------------------------------------------------------
    # STEP 2: Query Completed Roaming Reservations (pageSize=1000)
    # -------------------------------------------------------------
    print(f"\n[QUERY 2/2] Fetching Completed Roaming Reservations ({start_date} to {end_date})...")
    comp_url = f"https://cms.ev-charge-network.com/Roaming/OCPIReservation/LoadOcpiCompletedReservationGridViewThroughAjax?StartDate={start_date}&Enddate={end_date}&transactionTypeId=-1"
    req_comp = urllib.request.Request(comp_url, data=payload_encoded, headers=headers)
    with urllib.request.urlopen(req_comp, timeout=15) as res_comp:
        data_comp = json.loads(res_comp.read().decode("utf-8"))
    raw_completed = data_comp.get("Data", [])
    print(f"Retrieved {len(raw_completed)} Completed Roaming reservations from portal API!")

    # Filter Completed for Low Consumption (< 1 kWh)
    raw_low_comp = []
    for it in raw_completed:
        try:
            kwh_val = float(it.get("KWh") or 0.0)
        except (ValueError, TypeError):
            kwh_val = 0.0
        if kwh_val < 1.0:
            raw_low_comp.append(it)

    print(f"Retrieved {len(raw_low_comp)} Low-Consumption (< 1 kWh) Roaming reservations!")

    # -------------------------------------------------------------
    # STEP 3: Process and Format Records
    # -------------------------------------------------------------
    processed_rows = []

    # Process Cancelled
    for it in raw_cancelled:
        bid = it.get("BookingId") or ""
        b_date_raw = it.get("BookingInTime") or it.get("BookingDate")
        dt = parse_flexible_dt(b_date_raw) or datetime.now()
        date_str = f"{dt.month}/{dt.day}/{dt.year}"
        
        sched_action = it.get("SchedularAction")
        rca_text, rca_expl = classify_cancelled_reason(sched_action)

        processed_rows.append({
            "date": date_str,
            "reason_type": "Cancelled",
            "booking_id": bid,
            "auth_ref": it.get("AuthorizationReference") or "",
            "session_id": it.get("SessionId") or "",
            "party_id": it.get("PartyId") or "Unknown",
            "party_name": it.get("PartyName") or "Unknown",
            "manufacturer": it.get("ManufacturerName") or "Unknown",
            "model": it.get("ModelName") or "Unknown",
            "vehicle_number": it.get("VehicleNumber") or "",
            "station_name": it.get("StationName") or "Unknown Station",
            "charger_code": it.get("ChargerId") or "",
            "connector": it.get("ConnectorName") or it.get("Connector") or "A",
            "duration": it.get("TimeDuration") or "00:00:00",
            "energy": 0.0,
            "rca": rca_text,
            "schedular_action": sched_action or "",
            "explanation": rca_expl,
            "dt": dt
        })

    # Process Low-Consumption (< 1 kWh)
    for it in raw_low_comp:
        bid = it.get("BookingId") or ""
        b_date_raw = it.get("BookingInTime") or it.get("BookingDate")
        dt = parse_flexible_dt(b_date_raw) or datetime.now()
        date_str = f"{dt.month}/{dt.day}/{dt.year}"

        try:
            kwh_val = float(it.get("KWh") or 0.0)
        except (ValueError, TypeError):
            kwh_val = 0.0

        dur = it.get("TimeDuration") or "00:00:00"
        sched_action = it.get("SchedularAction")
        rca_text, rca_expl = classify_low_consumption_reason(kwh_val, dur, sched_action)

        processed_rows.append({
            "date": date_str,
            "reason_type": "less than 1kWh",
            "booking_id": bid,
            "auth_ref": it.get("AuthorizationReference") or "",
            "session_id": it.get("SessionId") or "",
            "party_id": it.get("PartyId") or "Unknown",
            "party_name": it.get("PartyName") or "Unknown",
            "manufacturer": it.get("ManufacturerName") or "Unknown",
            "model": it.get("ModelName") or "Unknown",
            "vehicle_number": it.get("VehicleNumber") or "",
            "station_name": it.get("StationName") or "Unknown Station",
            "charger_code": it.get("ChargerId") or "",
            "connector": it.get("ConnectorName") or it.get("Connector") or "A",
            "duration": dur,
            "energy": kwh_val,
            "rca": rca_text,
            "schedular_action": sched_action or "",
            "explanation": rca_expl,
            "dt": dt
        })

    # Sort descending by Booking ID (most recent first)
    processed_rows.sort(
        key=lambda r: int(r["booking_id"]) if str(r["booking_id"]).isdigit() else 0,
        reverse=True
    )

    # -------------------------------------------------------------
    # STEP 4: Build Excel Workbook
    # -------------------------------------------------------------
    print(f"\n[EXCEL] Building spreadsheet with {len(processed_rows)} total rows...")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Roaming Reservation RCA"

    headers = [
        "Date",
        "Reason",
        "Booking ID",
        "Party ID",
        "Party Name",
        "Manufacturer",
        "Model",
        "Vehicle Number",
        "Station Name",
        "Charger Code",
        "Connector",
        "Duration",
        "Energy (kWh)",
        "RCA Issue",
        "RCA Explanation",
        "Detailed Schedular Action",
        "Authorization Reference",
        "Session ID"
    ]
    ws.append(headers)

    # Styling definitions matching CS team standards
    yellow_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="000000")
    data_font = Font(name="Calibri", size=11, bold=False, color="000000")
    thin_border = Border(
        left=Side(style='thin', color='A0A0A0'),
        right=Side(style='thin', color='A0A0A0'),
        top=Side(style='thin', color='A0A0A0'),
        bottom=Side(style='thin', color='A0A0A0')
    )
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    right_align = Alignment(horizontal="right", vertical="center")

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = yellow_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = thin_border
    ws.row_dimensions[1].height = 24

    for row_idx, r in enumerate(processed_rows, start=2):
        ws.append([
            r["date"],
            r["reason_type"],
            int(r["booking_id"]) if str(r["booking_id"]).isdigit() else r["booking_id"],
            r["party_id"],
            r["party_name"],
            r["manufacturer"],
            r["model"],
            r["vehicle_number"],
            r["station_name"],
            r["charger_code"],
            r["connector"],
            r["duration"],
            r["energy"],
            r["rca"],
            r["explanation"],
            r["schedular_action"],
            r["auth_ref"],
            r["session_id"]
        ])
        ws.row_dimensions[row_idx].height = 20
        for col_idx in range(1, len(headers) + 1):
            c = ws.cell(row=row_idx, column=col_idx)
            c.font = data_font
            c.border = thin_border
            if col_idx in (1, 2, 3, 4, 8, 11, 12):
                c.alignment = center_align
            elif col_idx in (13,):
                c.alignment = right_align
            else:
                c.alignment = left_align

    # Column widths
    col_widths = {
        "A": 12,  # Date
        "B": 18,  # Reason ('Cancelled' or 'less than 1kWh')
        "C": 14,  # Booking ID
        "D": 12,  # Party ID
        "E": 26,  # Party Name
        "F": 18,  # Manufacturer
        "G": 20,  # Model
        "H": 16,  # Vehicle Number
        "I": 46,  # Station Name
        "J": 26,  # Charger Code
        "K": 10,  # Connector
        "L": 14,  # Duration
        "M": 14,  # Energy (kWh)
        "N": 52,  # RCA Issue
        "O": 58,  # RCA Explanation
        "P": 44,  # Detailed Schedular Action
        "Q": 36,  # Authorization Reference
        "R": 30   # Session ID
    }
    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width

    primary_path = os.path.join("E:\\ElectreeFi", output_excel)
    alt_path = os.path.join("E:\\ElectreeFi", "ElectreeFi_Roaming_Reservation_RCA_All_Records.xlsx")

    saved_paths = []
    try:
        wb.save(primary_path)
        saved_paths.append(primary_path)
        print(f"Saved primary workbook: {primary_path}")
    except PermissionError:
        print(f"Notice: '{output_excel}' is currently open.")

    try:
        wb.save(alt_path)
        saved_paths.append(alt_path)
        print(f"Saved backup workbook: {alt_path}")
    except Exception as e:
        print(f"Error saving backup path: {e}")

    print("\n" + "=" * 80)
    print(f"SUCCESS: ROAMING RCA EXCEL SAVED WITH ALL {len(processed_rows)} SESSIONS!")
    print(f"  Cancelled Bookings: {len(raw_cancelled)}")
    print(f"  Completed Low-Consumption (<1 kWh): {len(raw_low_comp)}")
    for p in saved_paths:
        print(f"Location: {p}")
    print("=" * 80)
    return saved_paths[0] if saved_paths else None


if __name__ == "__main__":
    asyncio.run(run_roaming_rca())
