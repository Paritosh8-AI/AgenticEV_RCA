import re
import json
import time
import threading

class HostRateLimiter:
    def __init__(self, min_interval_seconds: float = 0.20):
        self.min_interval = min_interval_seconds
        self.last_request_time: dict[str, float] = {}
        self.lock = threading.Lock()

    def acquire(self, host: str):
        with self.lock:
            now = time.time()
            last = self.last_request_time.get(host, 0.0)
            elapsed = now - last
            if elapsed < self.min_interval:
                sleep_time = self.min_interval - elapsed
                time.sleep(sleep_time)
            self.last_request_time[host] = time.time()

def classify_ocpi_text(text: str) -> tuple[bool, str]:
    clean = str(text or "").strip()
    if not clean or clean.lower() in ("null", "none", "n/a", "na", "-", "undefined", "noerror", "ok", "success"):
        return False, ""
    # Check for UUID
    if re.match(r"^[0-9a-fA-F]{8}-[0-9a-fA-F]{4}-[0-9a-fA-F]{4}-[0-9a-fA-F]{4}-[0-9a-fA-F]{12}$", clean):
        return False, clean
    # Check for pure hex/alphanumeric token without spaces
    if len(clean) >= 12 and re.match(r"^[0-9a-zA-Z_-]+$", clean) and not any(k in clean.lower() for k in ["error", "fail", "reject", "busy", "timeout", "offline", "abort", "fault"]):
        return False, clean
    lower = clean.lower()
    known = [
        "not connected", "disconnected", "busy", "occupied", "rejected", "timeout", "timed out",
        "invalid", "blocked", "unauthorized", "offline", "unreachable", "abort", "fault",
        "error", "failure", "fail", "stopped", "expired", "limit", "denied", "unknown", "refused",
        "start initiated"
    ]
    if any(k in lower for k in known):
        # "start initiated" is status info, not failure
        if "start initiated" in lower:
            return False, clean
        return True, clean
    words = clean.split()
    if len(words) >= 2 and all(w.isalnum() for w in words):
        return True, clean
    return False, clean

def parse_ocpi_start_session(raw_text: str) -> dict:
    text = str(raw_text or "").strip()
    clean_raw = text.split("??")[0].strip()
    if not clean_raw or clean_raw.upper() in ("N/A", "NA", "NULL", "NONE"):
        return {
            "is_na": True,
            "raw_response": text,
            "result": "NA",
            "ocpi_text": "",
            "is_real_reason": False,
            "failure_reason": ""
        }
    try:
        data = json.loads(clean_raw)
        item = data[0] if isinstance(data, list) and data else (data if isinstance(data, dict) else {})
        res_status = str(item.get("result") or item.get("status") or "").upper()
        messages = item.get("message") or []
        texts = []
        if isinstance(messages, list):
            for m in messages:
                if isinstance(m, dict) and m.get("text"):
                    texts.append(str(m["text"]).strip())
                elif isinstance(m, str):
                    texts.append(m.strip())
        combined_text = " | ".join(texts)
        is_reason, failure_text = classify_ocpi_text(combined_text)
        return {
            "is_na": False,
            "raw_response": clean_raw,
            "result": res_status,
            "ocpi_text": combined_text,
            "is_real_reason": is_reason,
            "failure_reason": failure_text if is_reason else ""
        }
    except Exception:
        # Check if text contains NA
        if "N/A" in clean_raw or "NA" in clean_raw:
            return {
                "is_na": True,
                "raw_response": text,
                "result": "NA",
                "ocpi_text": "",
                "is_real_reason": False,
                "failure_reason": ""
            }
        is_reason, failure_text = classify_ocpi_text(clean_raw)
        return {
            "is_na": False,
            "raw_response": clean_raw,
            "result": "UNKNOWN",
            "ocpi_text": clean_raw,
            "is_real_reason": is_reason,
            "failure_reason": failure_text if is_reason else ""
        }


def test_ocpi_start_session_na():
    assert parse_ocpi_start_session('N/A ?? "N/A"')["is_na"] is True
    assert parse_ocpi_start_session('NA')["is_na"] is True
    assert parse_ocpi_start_session('')["is_na"] is True

def test_ocpi_start_session_uuid():
    uuid_sample = '[{"result":"ACCEPTED","timeout":60,"message":[{"language":"en","text":"23b050ba-bebe-4c85-84f1-2fdffb7b02c3"}],"BookingId":0}] ?? "N/A"'
    p_uuid = parse_ocpi_start_session(uuid_sample)
    assert p_uuid["is_na"] is False
    assert p_uuid["result"] == "ACCEPTED"
    assert p_uuid["is_real_reason"] is False

def test_ocpi_start_session_error():
    error_sample = '[{"result":"REJECTED","timeout":60,"message":[{"language":"en","text":"EV not connected"}],"BookingId":0}] ?? "N/A"'
    p_err = parse_ocpi_start_session(error_sample)
    assert p_err["is_na"] is False
    assert p_err["result"] == "REJECTED"
    assert p_err["is_real_reason"] is True
    assert p_err["failure_reason"] == "EV not connected"

def test_filter_completed_low_consumption():
    def filter_completed_row(kwh: float, is_completed_sheet: bool):
        if is_completed_sheet and kwh >= 1.0:
            return False, "Discarded (Normal Session >= 1.0 kWh)"
        if is_completed_sheet:
            return True, "Completed (<1kWh)"
        return True, "Cancelled"

    assert filter_completed_row(18.5, True)[0] is False
    assert filter_completed_row(1.0, True)[0] is False
    assert filter_completed_row(0.99, True)[0] is True
    assert filter_completed_row(0.35, True)[1] == "Completed (<1kWh)"


def test_vehicle_model_aggregation_disambiguation():
    from src.rca.roaming_upload_analyzer import RoamingUploadAnalyzer
    analyzer = RoamingUploadAnalyzer()
    instances = [
        {
            "booking_id": "1001",
            "resolved_evse_id": "EVSE-01",
            "roaming_uid": "IOC-01",
            "party_id": "IOC",
            "station_name": "Delhi Central",
            "vehicle": "Nexon EV (DL01AB1234)",
            "vehicle_make": "Tata Motors",
            "vehicle_model": "Nexon EV",
            "manufacturer": "Tata Motors",
            "charger_model": "Delta UFC50",
            "charger_manufacturer": "DELTA",
            "root_cause": "BMS Handshake Timeout (~60s Protocol Window)",
            "fault_side": "VEHICLE SIDE",
            "kwh": 0.0,
            "session_category": "Cancelled"
        },
        {
            "booking_id": "1002",
            "resolved_evse_id": "EVSE-02",
            "roaming_uid": "IOC-02",
            "party_id": "IOC",
            "station_name": "Noida Hub",
            "vehicle": "XEV 9e (UP16XY9999)",
            "vehicle_make": "Mahindra",
            "vehicle_model": "XEV 9e",
            "manufacturer": "Mahindra",
            "charger_model": "Exicom Harmony",
            "charger_manufacturer": "EXICOM",
            "root_cause": "Charger Controller Initiation Timeout",
            "fault_side": "CHARGER SIDE",
            "kwh": 0.0,
            "session_category": "Cancelled"
        }
    ]

    aggs = analyzer.compute_aggregations(instances)
    v_stats = aggs["vehicle_model_stats"]
    assert len(v_stats) == 2
    tata = next(it for it in v_stats if it["mfg"] == "Tata Motors")
    assert tata["model"] == "Nexon EV"
    assert tata["vehicle_bms_faults"] == 1
    assert tata["dominant_side"] == "Vehicle Side"

    mahindra = next(it for it in v_stats if it["mfg"] == "Mahindra")
    assert mahindra["model"] == "XEV 9e"
    assert mahindra["charger_side"] == 1
    assert mahindra["vehicle_bms_faults"] == 0

    hw_stats = aggs["charger_hardware_stats"]
    assert len(hw_stats) == 2
    delta = next(it for it in hw_stats if it["mfg"] == "DELTA")
    assert delta["model"] == "Delta UFC50"
    exicom = next(it for it in hw_stats if it["mfg"] == "EXICOM")
    assert exicom["model"] == "Exicom Harmony"
    assert exicom["charger_hardware_faults"] == 1


def test_resolve_charger_hardware_lookup():
    from src.rca.roaming_upload_analyzer import RoamingUploadAnalyzer
    analyzer = RoamingUploadAnalyzer()

    # Test brand signature resolution
    res_tirex = analyzer.resolve_charger_hardware(station_name="IOCL Tirex Speedway 240kW")
    assert res_tirex["charger_manufacturer"] == "Tirex"
    assert "240" in res_tirex["charger_model"]

    res_exicom = analyzer.resolve_charger_hardware(station_name="Adhoc Exicom DC Fast Charging")
    assert res_exicom["charger_manufacturer"] == "EXICOM"

    res_delta = analyzer.resolve_charger_hardware(station_name="Highway Delta CCS240 Plaza")
    assert res_delta["charger_manufacturer"] == "DELTA"


def test_charger_model_and_oem_analysis_and_reports(tmp_path):
    import os
    from src.rca.roaming_upload_analyzer import RoamingUploadAnalyzer
    from docx import Document
    import openpyxl

    analyzer = RoamingUploadAnalyzer()
    instances = [
        {
            "booking_id": "2001",
            "resolved_evse_id": "EVSE-ZET-60",
            "roaming_uid": "IOC-ZET-01",
            "party_id": "IOC",
            "station_name": "IOCL Zetwerk Hub 60kW",
            "vehicle": "Nexon EV",
            "vehicle_make": "Tata Motors",
            "vehicle_model": "Nexon EV",
            "charger_manufacturer": "Zetwerk",
            "charger_model": "Zetwerk 60kW CCS2",
            "charger_capacity": 60.0,
            "root_cause": "Charger Controller Initiation Timeout",
            "fault_side": "CHARGER SIDE",
            "kwh": 0.0,
            "duration": 125,
            "in_time": "2026-10-01 10:00:00",
            "out_time": "2026-10-01 10:02:05",
            "connector_display": "Gun 1 (A)",
            "session_category": "Cancelled"
        },
        {
            "booking_id": "2002",
            "resolved_evse_id": "EVSE-QCH-180",
            "roaming_uid": "MPC-QCH-01",
            "party_id": "MPC",
            "station_name": "MPC Quench Speedway 180kW",
            "vehicle": "XEV 9e",
            "vehicle_make": "Mahindra",
            "vehicle_model": "XEV 9e",
            "charger_manufacturer": "Quench",
            "charger_model": "Quench 180kW Ultra-Fast",
            "charger_capacity": 180.0,
            "root_cause": "RemoteStart Rejected by Station",
            "fault_side": "CHARGER SIDE",
            "kwh": 0.0,
            "duration": 45,
            "in_time": "2026-10-01 11:00:00",
            "out_time": "2026-10-01 11:00:45",
            "connector_display": "Gun 2 (B)",
            "session_category": "Cancelled"
        },
        {
            "booking_id": "2003",
            "resolved_evse_id": "EVSE-CC-30",
            "roaming_uid": "VIN-CC-01",
            "party_id": "VIN",
            "station_name": "VinFast Chargecore Plaza",
            "vehicle": "VF8",
            "vehicle_make": "VinFast",
            "vehicle_model": "VF8",
            "charger_manufacturer": "Chargecore",
            "charger_model": "Chargecore 30kW DC Fast",
            "charger_capacity": 30.0,
            "root_cause": "Reservation Expired by CMS Scheduler",
            "fault_side": "USER / OPERATOR SIDE",
            "kwh": 0.0,
            "duration": 900,
            "in_time": "2026-10-01 12:00:00",
            "out_time": "2026-10-01 12:15:00",
            "connector_display": "Gun 1 (A)",
            "session_category": "Cancelled"
        }
    ]

    aggs = analyzer.compute_aggregations(instances)
    
    # 1. Verify aggregations
    assert "charger_mfg_stats" in aggs
    assert "charger_model_type_stats" in aggs
    assert "charger_power_class_stats" in aggs

    mfg_names = [m["mfg"] for m in aggs["charger_mfg_stats"]]
    assert "Zetwerk" in mfg_names
    assert "Quench" in mfg_names
    assert "Chargecore" in mfg_names

    models = [m["model"] for m in aggs["charger_model_type_stats"]]
    assert "Zetwerk 60kW CCS2" in models
    assert "Quench 180kW Ultra-Fast" in models

    pwr_classes = [p["power_class"] for p in aggs["charger_power_class_stats"]]
    assert "DC Ultra-Fast (150-360 kW)" in pwr_classes
    assert "DC Fast (50-120 kW)" in pwr_classes
    assert "DC Mid-Power (20-30 kW)" in pwr_classes

    # 2. Verify Word Report generation
    out_docx = str(tmp_path / "test_report.docx")
    analyzer.generate_word_report(instances, out_docx)
    assert os.path.exists(out_docx)
    doc = Document(out_docx)
    doc_text = " ".join([p.text for p in doc.paragraphs])
    assert "Charger Model Type, OEM & Hardware Reliability Analysis" in doc_text
    assert "5.1 Top Problematic Charger Manufacturers (OEMs)" in doc_text
    assert "5.2 Charger Hardware Model Types & Vulnerability Matrix" in doc_text
    assert "5.3 Power Rating & Charger Type Vulnerability Breakdown" in doc_text

    # 3. Verify Excel Report generation
    out_xlsx = str(tmp_path / "test_report.xlsx")
    analyzer.generate_excel_report(instances, out_xlsx)
    assert os.path.exists(out_xlsx)
    wb = openpyxl.load_workbook(out_xlsx)
    assert "Charger Model & OEM Analysis" in wb.sheetnames
    ws_hw = wb["Charger Model & OEM Analysis"]
    cell_values = [str(row[0]) for row in ws_hw.iter_rows(values_only=True) if row[0] is not None]
    assert any("1. CHARGER MANUFACTURERS" in v for v in cell_values)
    assert any("2. CHARGER HARDWARE MODEL TYPES" in v for v in cell_values)
    assert any("3. POWER RATING & CHARGER TYPE" in v for v in cell_values)


