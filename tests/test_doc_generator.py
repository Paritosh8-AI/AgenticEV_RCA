import os
import pytest
import docx
from src.reports.generate_word_report import build_rca_report

def test_build_rca_report(tmp_path):
    out_file = str(tmp_path / 'test_report.docx')
    res_path = build_rca_report(out_file)
    
    assert os.path.exists(res_path)
    assert os.path.getsize(res_path) > 10000
    
    # Verify Word structure
    doc = docx.Document(res_path)
    assert len(doc.tables) >= 5
    
    # Table 0: KPI Banner
    kpi_tbl = doc.tables[0]
    assert len(kpi_tbl.rows) == 1
    assert len(kpi_tbl.columns) == 3
    
    # Table 1: Cancelled Bookings occurrence table
    canc_tbl = doc.tables[1]
    assert len(canc_tbl.rows) > 1
    header_cells = [c.text.strip() for c in canc_tbl.rows[0].cells]
    assert 'RCA Reason' in header_cells
    assert 'Count' in header_cells
    
    # Table 2: Cancelled Fault Attribution
    attr_tbl = doc.tables[2]
    assert len(attr_tbl.rows) > 1
    attr_headers = [c.text.strip() for c in attr_tbl.rows[0].cells]
    assert 'Fault / Ownership' in attr_headers
